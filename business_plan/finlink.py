"""
The bridge between the financial model and the business plan.

The two documents used to hold their own numbers and were asked to agree. They
did not, and could not: the model was rebuilt around a contract book while the
plan still described a typed quarterly plan. So the plan no longer holds
numbers at all. It reads them from here, and this module reads them from the
recalculated workbook that `scripts/build_financial_model.py` produces.

That makes the workbook the single source of truth for every figure printed in
the plan. A number can be wrong, but the two files cannot disagree.

The workbook must have been recalculated before this is used, because openpyxl
writes formulas without cached values:

    python scripts/build_financial_model.py
    python scripts/recalc.py output/StartPad_Financial_Model.xlsx 1500
"""
import os
import sys

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from finmodel import plan as P                                   # noqa: E402

WORKBOOK = os.path.join(ROOT, "output", "StartPad_Financial_Model.xlsx")

YEARS = P.YEARS
MONTHS = P.MONTHS
YEAR_LABELS = P.YEAR_LABELS


# ------------------------------------------------------------ geometry
def ycol(y):
    """Column of the year-y total on a monthly sheet."""
    return 2 + (y - 1) * 13


def mcol(m):
    return ycol((m - 1) // 12 + 1) + (m - 1) % 12 + 1


def acol(y):
    """Column of year y on an annual sheet."""
    return 2 + y


class _Book:
    """The workbook, read once, addressed by sheet and row label."""

    def __init__(self, path=WORKBOOK):
        if not os.path.exists(path):
            raise SystemExit(
                f"{path} is missing. Build it first:\n"
                f"    python scripts/build_financial_model.py\n"
                f"    python scripts/recalc.py {path} 1500")
        self.wb = load_workbook(path, data_only=True)
        self._rows = {}
        probe = self.wb["Income Statement"].cell(
            self._row("Income Statement", "Total revenue"), acol(1)).value
        if not isinstance(probe, (int, float)):
            raise SystemExit(
                f"{path} has no cached values, so it has not been "
                f"recalculated. Run:\n"
                f"    python scripts/recalc.py {path} 1500")

    def _index(self, sheet):
        if sheet not in self._rows:
            ws = self.wb[sheet]
            d = {}
            for r in range(1, ws.max_row + 1):
                v = ws.cell(r, 1).value
                if isinstance(v, str) and v.strip():
                    d.setdefault(v.strip(), []).append(r)
            self._rows[sheet] = d
        return self._rows[sheet]

    def _row(self, sheet, label, nth=0):
        idx = self._index(sheet)
        if label not in idx:
            raise KeyError(f"{sheet!r} has no row labelled {label!r}")
        return idx[label][nth]

    def has(self, sheet, label):
        return label in self._index(sheet)

    def cell(self, sheet, label, col, nth=0):
        v = self.wb[sheet].cell(self._row(sheet, label, nth), col).value
        return v if isinstance(v, (int, float)) else 0.0

    def year(self, sheet, label, y, nth=0):
        """A year total on a monthly sheet."""
        return self.cell(sheet, label, ycol(y), nth)

    def years(self, sheet, label, nth=0):
        return [self.year(sheet, label, y, nth) for y in range(1, YEARS + 1)]

    def annual(self, sheet, label, y, nth=0):
        """A year column on an annual sheet."""
        return self.cell(sheet, label, acol(y), nth)

    def annuals(self, sheet, label, nth=0):
        return [self.annual(sheet, label, y, nth) for y in range(1, YEARS + 1)]

    def months(self, sheet, label, nth=0):
        return [self.cell(sheet, label, mcol(m), nth)
                for m in range(1, MONTHS + 1)]

    def assumption(self, label, col=3, nth=0):
        return self.cell("Assumption Setting", label, col, nth)

    def assumption_years(self, label, nth=0):
        return [self.assumption(label, 2 + y, nth) for y in range(1, YEARS + 1)]

    def text(self, sheet, label, col, nth=0):
        return self.wb[sheet].cell(self._row(sheet, label, nth), col).value


_BOOK = None


def book():
    global _BOOK
    if _BOOK is None:
        _BOOK = _Book()
    return _BOOK


def classify(number_format):
    """
    What kind of value a cell holds, taken from its own number format.

    Guessing the unit from the magnitude is how a count of 107 cohorts ends up
    printed as $107. The workbook already knows, so ask it.
    """
    f = (number_format or "").lower()
    # How many decimals the sheet itself shows. A cost per lead formatted to
    # two places is $4.66, and rounding it to $5 in the plan would be a
    # different number from the one in the workbook.
    dp = 0
    if "." in f:
        tail = f.split(".", 1)[1]
        dp = len(tail) - len(tail.lstrip("0"))
    if "%" in f:
        return "pct"
    if "x" in f and "$" not in f:
        return "mult"
    if "$" in f:
        return "money2" if dp else "money"
    if '"days"' in f:
        return "days"
    return f"num{dp}" if dp else "num"


def _fmt_of(ws, row, cols):
    """
    The kind of a row, from the first column on it that actually holds a
    number. Probing one fixed column misreads a line whose first year is
    blank, which is exactly the case for a growth rate.
    """
    if isinstance(cols, int):
        cols = [cols]
    for c in cols:
        if isinstance(ws.cell(row, c).value, (int, float)):
            return classify(ws.cell(row, c).number_format)
    return classify(ws.cell(row, cols[0]).number_format)


# ------------------------------------------------------- the statements
def income_statement():
    """Every line of the income statement, in the order it is presented."""
    b = book()
    ws = b.wb["Income Statement"]
    out = []
    for r in range(5, ws.max_row + 1):
        label = ws.cell(r, 1).value
        if not isinstance(label, str) or not label.strip():
            continue
        vals = [ws.cell(r, acol(y)).value for y in range(1, YEARS + 1)]
        kind = _fmt_of(ws, r, [acol(y) for y in range(1, YEARS + 1)])
        if all(not isinstance(v, (int, float)) for v in vals):
            out.append(dict(label=label.strip(), values=None, heading=True,
                            kind=kind, indent=label.startswith(" ")))
            continue
        out.append(dict(label=label.strip(), kind=kind,
                        values=[v if isinstance(v, (int, float)) else 0.0
                                for v in vals],
                        heading=False, indent=label.startswith(" ")))
    return out


# The balance sheet carries the opening position in its first column, so its
# year columns start one to the right of every other annual sheet.
BS_OPENING_COL = acol(1)


def bscol(y):
    return acol(y) + 1


def balance_sheet():
    """Each line at each year end, plus the opening position it starts from."""
    b = book()
    ws = b.wb["Balance Sheet"]
    out = []
    for r in range(5, ws.max_row + 1):
        label = ws.cell(r, 1).value
        if not isinstance(label, str) or not label.strip():
            continue
        opening = ws.cell(r, BS_OPENING_COL).value
        vals = [ws.cell(r, bscol(y)).value for y in range(1, YEARS + 1)]
        kind = _fmt_of(ws, r, [bscol(y) for y in range(1, YEARS + 1)])
        out.append(dict(label=label.strip(), kind=kind,
                        opening=opening if isinstance(opening, (int, float))
                        else None,
                        values=[v if isinstance(v, (int, float)) else 0.0
                                for v in vals],
                        heading=all(not isinstance(v, (int, float))
                                    for v in vals),
                        indent=label.startswith(" ")))
    return out


def cash_flow():
    """The annual cash flow, taken from the monthly sheet's year totals."""
    b = book()
    ws = b.wb["Cash Flow Statement"]
    out = []
    for r in range(5, ws.max_row + 1):
        label = ws.cell(r, 1).value
        if not isinstance(label, str) or not label.strip():
            continue
        if len(label) > 90 or label.strip() == "Line item":
            continue
        vals = [ws.cell(r, ycol(y)).value for y in range(1, YEARS + 1)]
        kind = _fmt_of(ws, r, [ycol(y) for y in range(1, YEARS + 1)])
        out.append(dict(label=label.strip(), kind=kind,
                        values=[v if isinstance(v, (int, float)) else 0.0
                                for v in vals],
                        heading=all(not isinstance(v, (int, float))
                                    for v in vals),
                        indent=label.startswith(" ")))
    return out


# ------------------------------------------------------------ headlines
def headline():
    b = book()
    rev = b.annuals("Income Statement", "Total revenue")
    cos = b.annuals("Income Statement", "Cost of revenue")
    gp = b.annuals("Income Statement", "Gross profit")
    opex = b.annuals("Income Statement", "Total operating expenses")
    ebitda = b.annuals("Income Statement", "EBITDA")
    tax = b.annuals("Income Statement", "Corporate income tax")
    net = b.annuals("Income Statement", "Net profit for the year")
    total_exp = b.annuals("Income Statement",
                          "Total expenses, direct and operating")
    cash = b.years("Cash Flow Statement", "Cash at the end of the month")
    return [dict(year=y + 1,
                 label=YEAR_LABELS[y],
                 revenue=rev[y], cost_of_revenue=cos[y], gross_profit=gp[y],
                 gross_margin=(gp[y] / rev[y]) if rev[y] else 0.0,
                 opex=opex[y], ebitda=ebitda[y],
                 ebitda_margin=(ebitda[y] / rev[y]) if rev[y] else 0.0,
                 tax=tax[y], net=net[y],
                 net_margin=(net[y] / rev[y]) if rev[y] else 0.0,
                 total_expense=total_exp[y], cash_end=cash[y])
            for y in range(YEARS)]


def first_profitable_year():
    for h in headline():
        if h["net"] > 0:
            return h["year"]
    return None


# --------------------------------------------------------------- revenue
REVENUE_LINES = [
    ("University cohort programmes", "University cohorts, Egypt", "Education"),
    ("Faculty licences", "Faculty licences", "Education"),
    ("School cohorts", "School cohorts", "Education"),
    ("School licences", "School licences", "Education"),
    ("Technical and vocational cohorts", "Technical and vocational cohorts",
     "Education"),
    ("National initiative seats", "National initiative seats", "Public"),
    ("Train the trainer cohorts", "Train the trainer cohorts", "Public"),
    ("NGO and donor programmes", "NGO and donor programmes", "Public"),
    ("Sponsored competitions", "Competitions, sponsored", "Commercial"),
    ("Corporate and CSR programmes", "Corporate and CSR programmes",
     "Commercial"),
    ("Programme and incubator seats", "Programme and incubator seats",
     "Education"),
    ("Mission Zero sponsorship", "Mission Zero seasons carrying a sponsor",
     "Commercial"),
]

_DUP = {"Faculty licences", "School cohorts", "School licences",
        "Technical and vocational cohorts", "National initiative seats",
        "Train the trainer cohorts", "NGO and donor programmes",
        "Corporate and CSR programmes", "Programme and incubator seats"}


def revenue_lines():
    """Each revenue stream: its volume, its unit price and its revenue."""
    b = book()
    out = []
    for stream, vol_label, family in REVENUE_LINES:
        nth = 1 if stream in _DUP else 0
        rev = b.years("Revenue Streams", stream, nth)
        vol = b.years("Revenue Streams", vol_label)
        out.append(dict(stream=stream, family=family, unit=vol_label,
                        revenue=rev, volume=vol,
                        price=[(r / v if v else 0.0)
                               for r, v in zip(rev, vol)]))
    subs = b.years("Revenue Streams", "Subscription revenue recognised")
    out.append(dict(stream="Founder subscriptions", family="Direct",
                    unit="Active paid accounts",
                    revenue=subs,
                    volume=b.years("Revenue Streams",
                                   "Active paid subscribers"),
                    price=[0.0] * YEARS))
    return out


def revenue_by_family():
    fams = {}
    for line in revenue_lines():
        acc = fams.setdefault(line["family"], [0.0] * YEARS)
        for i, v in enumerate(line["revenue"]):
            acc[i] += v
    return fams


# ----------------------------------------------------------------- cost
def expense_lines():
    """Direct cost and operating cost, line by line, as the model builds it."""
    b = book()
    ws = b.wb["Expenses"]
    direct, operating = [], []
    bucket = None
    for r in range(5, ws.max_row + 1):
        label = ws.cell(r, 1).value
        if not isinstance(label, str) or not label.strip():
            continue
        label = label.strip()
        if label.startswith("DIRECT COST"):
            bucket = direct
            continue
        if label.startswith("OPERATING COST"):
            bucket = operating
            continue
        if label.startswith("Total ") or label.startswith("Gross ") \
                or label.startswith("EBIT") or label.startswith("Year one") \
                or label.startswith("Cash operating") \
                or label.startswith("Operating cost before"):
            continue
        vals = [ws.cell(r, ycol(y)).value for y in range(1, YEARS + 1)]
        if bucket is None or not any(isinstance(v, (int, float)) for v in vals):
            continue
        bucket.append(dict(label=label,
                           values=[v if isinstance(v, (int, float)) else 0.0
                                   for v in vals]))
    return dict(direct=direct, operating=operating,
                direct_total=b.years("Expenses", "Total direct cost"),
                operating_total=b.years("Expenses", "Total operating cost"),
                total=b.years("Expenses",
                              "Total expenses, direct and operating"))


# ------------------------------------------------------------- the funnel
def funnel():
    b = book()
    return dict(
        online_cost=b.years("Acquisitions", "Online marketing cost"),
        offline_cost=b.years("Acquisitions", "Offline marketing cost"),
        marketing=b.years("Acquisitions", "Total marketing cost"),
        clicks=b.years("Acquisitions", "Total clicks"),
        online_leads=b.years("Acquisitions", "Total online leads"),
        offline_leads=b.years("Acquisitions", "Offline institutional leads"),
        leads=b.years("Acquisitions", "Total institutional leads"),
        signed=b.years("Acquisitions", "Total contracts signed"),
        mobilised=b.years("Acquisitions", "Total contracts mobilised"),
        registered=b.years("Acquisitions", "Total new registered accounts"),
        new_paid=b.years("Acquisitions", "Stage 3, new paid accounts"),
        active_paid=b.years("Acquisitions", "Active paid accounts"),
        cac_b2c=b.years("Acquisitions", "Cost per paid account"),
        cac_inst=b.years("Acquisitions", "Cost per institutional contract won"),
    )


def buyers():
    """Each buyer type, with its conversion chain and its clock."""
    b = book()
    out = []
    for i, (key, label, share, l2m, m2p, p2w, cycle, mob, why) in enumerate(
            P.BUYERS):
        out.append(dict(
            key=key, label=label,
            share=b.assumption(label, 3), lead_to_meeting=b.assumption(label, 4),
            meeting_to_proposal=b.assumption(label, 5),
            proposal_to_win=b.assumption(label, 6),
            cycle=int(b.assumption(label, 7)),
            mobilisation=int(b.assumption(label, 8)),
            why=why,
            leads=b.years("Acquisitions", label),
            meetings=b.years("Acquisitions", "Meetings held", i),
            proposals=b.years("Acquisitions", "Proposals, pilots and bids", i),
            signed=b.years("Acquisitions",
                           "Contracts signed, after the sales cycle", i),
        ))
    return out


def engines():
    """Each engine: what a won contract turns into, and what it delivered."""
    b = book()
    out = []
    for (key, label, buyer, share, units, life, unit, gate, why) in P.ENGINES:
        out.append(dict(
            key=key, label=label, buyer=buyer, unit=unit, why=why,
            share=b.assumption(label, 3), units=b.assumption(label, 4),
            life=int(b.assumption(label, 5)), gate=int(b.assumption(label, 6)),
            volume=b.years("Revenue Streams", label),
            live=b.years("Acquisitions", label + ", contracts live"),
        ))
    return out


def online_channels():
    b = book()
    out = []
    for key, label, budgets, cpc, org, ref, s1 in P.ONLINE_CHANNELS:
        out.append(dict(key=key, label=label,
                        budget=b.assumption_years(label),
                        cpc=b.assumption(label, 8),
                        organic=b.assumption(label, 9),
                        referral=b.assumption(label, 10),
                        click_to_lead=b.assumption(label, 11)))
    return out


def offline_channels():
    b = book()
    out = []
    for key, label, acts, cost, signups, leads, why in P.OFFLINE_CHANNELS:
        out.append(dict(key=key, label=label, why=why,
                        activities=b.assumption_years(label),
                        cost=b.assumption(label, 8),
                        signups=b.assumption(label, 9),
                        leads=b.assumption(label, 10)))
    return out


# ------------------------------------------------------------- capacity
def capacity():
    b = book()
    return dict(
        activity_use=b.years("GTM Capacity and Timing", "Activity utilisation"),
        meeting_use=b.years("GTM Capacity and Timing", "Meeting utilisation"),
        activity_short=b.years("GTM Capacity and Timing",
                               "Activity the plan asks for beyond what the "
                               "team can run"),
        meeting_short=b.years("GTM Capacity and Timing",
                              "Meetings asked for beyond capacity"),
        facilitators=b.years("GTM Capacity and Timing",
                             "Certified facilitators in the pool"),
        partners=b.years("GTM Capacity and Timing",
                         "Certified delivery partner organisations"),
        fac_days=b.years("GTM Capacity and Timing",
                         "Facilitator days the cohort plan needs"),
        live=b.years("GTM Capacity and Timing",
                     "Contracts live across all engines"),
    )


def establishment():
    """The payroll, by function and by year, with the cap it is held to."""
    b = book()
    functions = list(dict.fromkeys(fn for fn, *_ in P.ROLES))
    by_function = {fn: b.years("Payroll", fn) for fn in functions}
    roles = []
    for function, role, start, gross, counts in P.ROLES:
        roles.append(dict(function=function, role=role, start=start,
                          gross=list(gross), counts=list(counts)))
    return dict(roles=roles, by_function=by_function,
                headcount=b.years("Payroll", "Total headcount"),
                payroll=b.years("Payroll", "Total payroll"),
                cap=list(P.HEADCOUNT_CAP))


# ------------------------------------------------------------- analysis
def indicators():
    b = book()
    ws = b.wb["Financial Indicators"]
    out = []
    for r in range(5, ws.max_row + 1):
        label = ws.cell(r, 1).value
        if not isinstance(label, str) or not label.strip():
            continue
        vals = [ws.cell(r, acol(y)).value for y in range(1, YEARS + 1)]
        if not any(isinstance(v, (int, float)) for v in vals):
            continue
        # A cell can legitimately read "n/a" in year one, where there is no
        # prior year to compare with. Keep the text rather than turning it
        # into a nil that would print as a real zero.
        out.append(dict(label=label.strip(), note=ws.cell(r, 2).value or "",
                        kind=_fmt_of(ws, r, [acol(y) for y in range(1, YEARS + 1)]),
                        values=[v if isinstance(v, (int, float)) else
                                (v if isinstance(v, str) else 0.0)
                                for v in vals]))
    return out


def cash_profile():
    """Monthly cash, and the trough the funding ask has to cover."""
    b = book()
    cash = b.months("Cash Flow Statement", "Cash at the end of the month")
    trough = min(cash)
    return dict(monthly=cash, trough=trough, trough_month=cash.index(trough) + 1,
                year_end=b.years("Cash Flow Statement",
                                 "Cash at the end of the month"),
                seed=b.assumption("Seed round raised"),
                seed_month=int(b.assumption("Month the seed round lands")),
                opening=b.assumption("Opening cash"),
                buffer=b.assumption("Minimum cash buffer to hold"))


def working_capital():
    b = book()
    return dict(dso=b.assumption("Days sales outstanding, institutional"),
                dso_sub=b.assumption("Days sales outstanding, subscriptions"),
                dpo=b.assumption("Days payable outstanding"),
                accrued=b.assumption("Accrued payroll, days outstanding"),
                advance=b.assumption("Institutional contracts invoiced in "
                                     "advance"))


def tax_basis():
    b = book()
    return dict(rate=b.assumption("Corporate income tax rate"),
                month=int(b.assumption("Month of the following year in which "
                                       "tax is paid")),
                opening_losses=b.assumption("Tax losses carried forward at "
                                            "inception"))


def prices():
    b = book()
    out = []
    for key, (label, unit, vals, why) in P.PRICES.items():
        out.append(dict(key=key, label=label, unit=unit, why=why,
                        values=b.assumption_years(label)))
    return out


def traction():
    b = book()
    return [dict(key=key, label=label, value=b.assumption(label), unit=unit,
                 note=note)
            for key, label, value, unit, note in P.TRACTION]


def pace():
    b = book()
    return b.assumption_years("Delivery pace against the full contract rate")


def summary_check():
    """What the workbook says about itself, so the plan can quote it."""
    b = book()
    return dict(
        year_one_expense=b.annual("Income Statement",
                                  "Total expenses, direct and operating", 1),
        balance_check=[b.cell("Balance Sheet", lbl, bscol(y))
                       for y in range(1, YEARS + 1)
                       for lbl in [next(k for k in b._index("Balance Sheet")
                                        if k.startswith("CHECK"))]],
    )
