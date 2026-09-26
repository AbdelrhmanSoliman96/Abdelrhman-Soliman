import re, pathlib
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x0F, 0x2C, 0x4A)
MIDBLUE = RGBColor(0x1C, 0x4A, 0x73)
GREY = RGBColor(0x55, 0x5F, 0x6B)
SANS = "Calibri"
SERIF = "Cambria"

import sys
SRC = pathlib.Path(sys.argv[1])
OUT = pathlib.Path(sys.argv[2])
FOOTER_LABEL = sys.argv[3]
COVER = sys.argv[4] if len(sys.argv) > 4 else None

doc = Document()

# ---------- page setup ----------
for s in doc.sections:
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0)
    s.top_margin = Cm(2.0); s.bottom_margin = Cm(2.0)

if COVER == 'agreement':
    import coverpage
    doc.sections[0].different_first_page_header_footer = True
    coverpage.build(
        doc, doc.sections[0],
        consultant   = "ABDELRHMAN SOLIMAN",
        client_legal = "[ INSERT FULL REGISTERED LEGAL NAME ]",
        client_trade = "BTAKKA",
        date_line    = "Dated  [ ____ ]  [ __________ ]  2026",
        proposal_ref = 'Relates to the Proposal "Investment Readiness & Investment Round\nManagement", version 1.0, dated 13 September 2026, annexed hereto',
        version      = "DOCUMENT 2 OF 2   ·   VERSION 1.0   ·   FOR CLIENT REVIEW   ·   UNEXECUTED",
        governing_law= "Arab Republic of Egypt",
    )

st = doc.styles['Normal']
st.font.name = SERIF; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), SERIF)
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.12

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexcolor)
    tcPr.append(sh)

def keep_with_next(p):
    p.paragraph_format.keep_with_next = True

def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader'); el.set(qn('w:val'), 'true'); trPr.append(el)

# ---------- inline markdown ----------
TOKEN = re.compile(r'(\*\*.+?\*\*|(?<!\*)\*(?!\*).+?(?<!\*)\*(?!\*)|`[^`]+`|\[[^\]]+\]\([^)]+\))', re.S)

def add_runs(par, text, base_bold=False, size=None, color=None, font=None):
    text = text.replace('<br>', ' ').replace('&amp;', '&').replace('&nbsp;', ' ')
    for part in TOKEN.split(text):
        if not part:
            continue
        bold, italic, mono = base_bold, False, False
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            part, bold = part[2:-2], True
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            part, italic = part[1:-1], True
        elif part.startswith('`') and part.endswith('`'):
            part, mono = part[1:-1], True
        elif part.startswith('['):
            m = re.match(r'\[([^\]]+)\]\(([^)]+)\)', part)
            if m:
                part = m.group(1)
                italic = False
        r = par.add_run(part)
        r.bold = bold; r.italic = italic
        r.font.name = "Consolas" if mono else (font or SERIF)
        if size: r.font.size = size
        if color: r.font.color.rgb = color
    return par

def heading(text, level, first_h1=[True]):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    if level == 1:
        if not first_h1[0]:
            p.add_run().add_break(WD_BREAK.PAGE)
        first_h1[0] = False
        pf.space_before, pf.space_after = Pt(0), Pt(10)
        add_runs(p, text, base_bold=True, size=Pt(19), color=NAVY, font=SANS)
        pbdr = OxmlElement('w:pBdr'); bot = OxmlElement('w:bottom')
        bot.set(qn('w:val'), 'single'); bot.set(qn('w:sz'), '12')
        bot.set(qn('w:space'), '4'); bot.set(qn('w:color'), '0F2C4A')
        pbdr.append(bot); p._p.get_or_add_pPr().append(pbdr)
    elif level == 2:
        pf.space_before, pf.space_after = Pt(16), Pt(5)
        add_runs(p, text, base_bold=True, size=Pt(13.5), color=NAVY, font=SANS)
    elif level == 3:
        pf.space_before, pf.space_after = Pt(12), Pt(4)
        add_runs(p, text, base_bold=True, size=Pt(11.5), color=MIDBLUE, font=SANS)
    else:
        pf.space_before, pf.space_after = Pt(10), Pt(3)
        add_runs(p, text, base_bold=True, size=Pt(10.5), color=MIDBLUE, font=SANS)
    keep_with_next(p)

def add_table(rows):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ''
        p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
        add_runs(p, h, base_bold=True, size=Pt(8.5), color=RGBColor(0xFF,0xFF,0xFF), font=SANS)
        shade(c, '0F2C4A')
    set_repeat_header(t.rows[0])
    for ri, row in enumerate(body):
        cells = t.add_row().cells
        for i in range(len(header)):
            val = row[i] if i < len(row) else ''
            c = cells[i]; c.text = ''
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
            add_runs(p, val, size=Pt(8.5))
            if ri % 2 == 1:
                shade(c, 'F2F5F8')
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def add_quote(lines):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.5); pf.right_indent = Cm(0.3)
    pf.space_before = Pt(7); pf.space_after = Pt(9)
    add_runs(p, ' '.join(lines), size=Pt(9.5), color=RGBColor(0x1C,0x3A,0x55))
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    for side, sz in (('left','18'),):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'), sz)
        el.set(qn('w:space'),'8'); el.set(qn('w:color'),'0F2C4A')
        pbdr.append(el)
    pPr.append(pbdr)
    for r in p.runs: r.italic = False
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'),'clear')
    shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),'EEF3F8')
    pPr.append(shd)

def add_rule():
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(6)
    pbdr = OxmlElement('w:pBdr'); bot = OxmlElement('w:bottom')
    bot.set(qn('w:val'),'single'); bot.set(qn('w:sz'),'6')
    bot.set(qn('w:space'),'1'); bot.set(qn('w:color'),'C2CCD6')
    pbdr.append(bot); p._p.get_or_add_pPr().append(pbdr)

# ---------- parse ----------
lines = SRC.read_text(encoding='utf-8').split('\n')
if COVER == 'agreement':
    # the designed title page replaces the markdown cover; resume at Document Control
    for n, ln in enumerate(lines):
        if ln.strip() == '### Document Control':
            lines = lines[n:]
            break
    else:
        raise SystemExit('cover mode: "### Document Control" marker not found')
i = 0
para_buf, quote_buf = [], []

def flush_para():
    global para_buf
    if para_buf:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        add_runs(p, ' '.join(para_buf))
        para_buf = []

def flush_quote():
    global quote_buf
    if quote_buf:
        add_quote(quote_buf); quote_buf = []

def split_row(line):
    s = line.strip()
    if s.startswith('|'): s = s[1:]
    if s.endswith('|'): s = s[:-1]
    return [c.strip() for c in s.split('|')]

while i < len(lines):
    line = lines[i]; s = line.strip()

    if not s:
        flush_para(); flush_quote(); i += 1; continue

    if s in ('<br>', '<br/>'):
        i += 1; continue

    if s.startswith('>'):
        flush_para(); quote_buf.append(s.lstrip('> ').strip()); i += 1; continue
    flush_quote()

    if re.fullmatch(r'-{3,}', s):
        flush_para(); add_rule(); i += 1; continue

    m = re.match(r'^(#{1,4})\s+(.*)$', s)
    if m:
        flush_para(); heading(m.group(2).strip(), len(m.group(1))); i += 1; continue

    # table
    if s.startswith('|') and i + 1 < len(lines) and re.match(r'^\s*\|[\s:|-]+\|\s*$', lines[i+1]):
        flush_para()
        rows = [split_row(lines[i])]
        j = i + 2
        while j < len(lines) and lines[j].strip().startswith('|'):
            rows.append(split_row(lines[j])); j += 1
        add_table(rows); i = j; continue

    # lists
    m = re.match(r'^(\s*)([-*]|\d+\.)\s+(.*)$', line)
    if m:
        flush_para()
        indent = len(m.group(1))
        ordered = not m.group(2) in ('-', '*')
        p = doc.add_paragraph(style='List Number' if ordered else 'List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Cm(0.7 + 0.5 * (indent // 2))
        add_runs(p, m.group(3))
        i += 1; continue

    para_buf.append(s); i += 1

flush_para(); flush_quote()

# ---------- footer ----------
from docx.oxml import OxmlElement as OE
sec = doc.sections[0]
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run(f"{FOOTER_LABEL}  \u00b7  Btakka \u00d7 Abdelrhman Soliman  \u00b7  v1.0  \u00b7  Page ")
r.font.size = Pt(7.5); r.font.color.rgb = GREY; r.font.name = SANS
fld = OE('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE')
fp._p.append(fld)

doc.save(OUT)
print("Saved:", OUT)
