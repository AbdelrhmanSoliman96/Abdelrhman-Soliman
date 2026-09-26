"""Designed title page for the Consulting Agreement DOCX."""
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY  = RGBColor(0x0F, 0x2C, 0x4A)
SLATE = RGBColor(0x55, 0x5F, 0x6B)
BRASS = RGBColor(0x8A, 0x61, 0x15)
SANS  = "Calibri"
SERIF = "Cambria"

def _p(doc, space_before=0, space_after=0, align=WD_ALIGN_PARAGRAPH.CENTER):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.0
    return p

def _run(p, text, size, color, font=SERIF, bold=False, caps_space=None, italic=False):
    r = p.add_run(text)
    r.font.size = Pt(size); r.font.color.rgb = color
    r.font.name = font; r.bold = bold; r.italic = italic
    r.element.rPr.rFonts.set(qn('w:eastAsia'), font)
    if caps_space is not None:
        sp = OxmlElement('w:spacing'); sp.set(qn('w:val'), str(int(caps_space * 20)))
        r.element.rPr.append(sp)
    return r

def _rule(doc, color='0F2C4A', size='8', space_before=0, space_after=0, width_pct=None):
    p = _p(doc, space_before, space_after)
    if width_pct:
        p.paragraph_format.left_indent = Cm(17.0 * (1 - width_pct / 100) / 2)
        p.paragraph_format.right_indent = Cm(17.0 * (1 - width_pct / 100) / 2)
    pbdr = OxmlElement('w:pBdr'); bot = OxmlElement('w:bottom')
    bot.set(qn('w:val'), 'single'); bot.set(qn('w:sz'), size)
    bot.set(qn('w:space'), '0'); bot.set(qn('w:color'), color)
    pbdr.append(bot); p._p.get_or_add_pPr().append(pbdr)
    return p

def build(doc, section, *, consultant, client_legal, client_trade,
          date_line, proposal_ref, version, governing_law):
    # ── eyebrow ────────────────────────────────────────────────
    p = _p(doc, 0, 4)
    _run(p, "P R I V A T E   &   C O N F I D E N T I A L", 8.5, BRASS, SANS, bold=True, caps_space=1.2)

    # breathing room before the title block
    for _ in range(6):
        _p(doc, 0, 0).add_run("")

    # ── title block ────────────────────────────────────────────
    _rule(doc, '0F2C4A', '12', 0, 16, width_pct=64)

    p = _p(doc, 0, 6)
    _run(p, "CONSULTING AGREEMENT", 27, NAVY, SANS, bold=True, caps_space=2.6)

    p = _p(doc, 4, 4)
    _run(p, "Investment Readiness & Investment Round Management", 13.5, SLATE, SERIF, italic=True)

    _rule(doc, '0F2C4A', '12', 14, 0, width_pct=64)

    for _ in range(4):
        _p(doc, 0, 0).add_run("")

    # ── parties ────────────────────────────────────────────────
    p = _p(doc, 0, 7); _run(p, "B E T W E E N", 8.5, BRASS, SANS, bold=True, caps_space=1.6)
    p = _p(doc, 0, 3); _run(p, consultant, 14, NAVY, SERIF, bold=True)
    p = _p(doc, 0, 16); _run(p, '"the Consultant"', 10, SLATE, SERIF, italic=True)

    p = _p(doc, 0, 7); _run(p, "A N D", 8.5, BRASS, SANS, bold=True, caps_space=1.6)
    p = _p(doc, 0, 2); _run(p, client_legal, 14, NAVY, SERIF, bold=True)
    p = _p(doc, 0, 3); _run(p, f"trading as {client_trade}", 11.5, NAVY, SERIF)
    p = _p(doc, 0, 0); _run(p, '"the Client"', 10, SLATE, SERIF, italic=True)

    for _ in range(4):
        _p(doc, 0, 0).add_run("")

    # ── date ───────────────────────────────────────────────────
    _rule(doc, 'C2CCD6', '6', 0, 12, width_pct=22)
    p = _p(doc, 0, 0); _run(p, date_line, 12, NAVY, SERIF, bold=True)
    _rule(doc, 'C2CCD6', '6', 12, 0, width_pct=22)

    # ── footer block ──────────────────────────────────────────
    # Written into Word's first-page footer so it is anchored to the bottom
    # margin by the layout engine, not by a guessed run of blank space.
    fp = section.first_page_footer
    fp.is_linked_to_previous = False
    for para in list(fp.paragraphs)[1:]:
        para._p.getparent().remove(para._p)
    base = fp.paragraphs[0]
    base.alignment = WD_ALIGN_PARAGRAPH.CENTER
    base.paragraph_format.space_after = Pt(3)
    base.paragraph_format.line_spacing = 1.15
    _run(base, proposal_ref.replace("\n", " "), 9, SLATE, SERIF, italic=True)

    for text, size, font in ((version, 8.5, SANS),
                             (f"Governed by the laws of the {governing_law}", 8.5, SANS)):
        q = fp.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q.paragraph_format.space_before = Pt(0)
        q.paragraph_format.space_after = Pt(3)
        _run(q, text, size, SLATE, font)

    # ── page break ─────────────────────────────────────────────
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
