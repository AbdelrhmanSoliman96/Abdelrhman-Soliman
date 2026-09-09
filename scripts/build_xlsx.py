"""Build Startpad — Mission x Tool Matrix.xlsx"""
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from startpad.missions import ALWAYS_AVAILABLE, MAPPING, MISSIONS  # noqa: E402
from startpad.resources import BY_MISSION, FOUNDATION, FUNDRAISING  # noqa: E402
from startpad.tools import CATEGORY_ORDER, TOOL_BY_NAME, TOOLS  # noqa: E402
from scripts._determinism import FIXED_TIMESTAMP, normalize_zip  # noqa: E402

INK = "1F2933"
ACCENT = "0F5C4B"
CORE = "0F5C4B"
SUPP = "1D6FA3"
STRETCH = "7A6A55"
HEAD_FILL = PatternFill("solid", fgColor="0F5C4B")
SUB_FILL = PatternFill("solid", fgColor="E8F0ED")
ZEBRA = PatternFill("solid", fgColor="F7F9F8")
THIN = Side(style="thin", color="D6DEDA")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

ROLE_FILL = {
    "Core": PatternFill("solid", fgColor="D7EBE3"),
    "Supporting": PatternFill("solid", fgColor="DCE9F4"),
    "Stretch": PatternFill("solid", fgColor="F0EAE0"),
}
ROLE_FONT = {
    "Core": Font(color=CORE, bold=True, size=10),
    "Supporting": Font(color=SUPP, bold=True, size=10),
    "Stretch": Font(color=STRETCH, bold=True, size=10),
}
ROLE_LABEL = {"core": "Core", "supporting": "Supporting", "stretch": "Stretch"}


def title_block(ws, title, subtitle, width):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    c = ws.cell(row=1, column=1, value=title)
    c.font = Font(size=16, bold=True, color="FFFFFF")
    c.fill = HEAD_FILL
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 30
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    c = ws.cell(row=2, column=2 - 1, value=subtitle)
    c.font = Font(size=10, italic=True, color=INK)
    c.fill = SUB_FILL
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[2].height = 20


def header_row(ws, row, headers, widths):
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor="2E6F5E")
        c.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
        c.border = BOX
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 28
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def add_table(ws, name, first_row, last_row, last_col):
    ref = f"A{first_row}:{get_column_letter(last_col)}{last_row}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(
        name="TableStyleLight1", showRowStripes=True, showColumnStripes=False
    )
    ws.add_table(t)


# ---------------------------------------------------------------- README
def sheet_readme(wb):
    ws = wb.create_sheet("How to use")
    ws.sheet_view.showGridLines = False
    title_block(ws, "Startpad — Mission × Tool Matrix",
                "Which of the 122 tools belong inside each of the 15 missions, and why.", 6)
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 96

    rows = [
        ("", ""),
        ("WHAT THIS IS", ""),
        ("", "A tool-to-mission assignment for the Startpad platform. Every tool in the Tool "
             "Library is placed in one of three ways: inside a mission as Core, inside a "
             "mission as Supporting/Stretch, or on an always-available shelf outside the "
             "mission flow. Nothing is left unplaced."),
        ("", ""),
        ("THE THREE ROLES", ""),
        ("Core",
         "The founder cannot write a good answer to this mission's questions without it. "
         "Surface it in the mission workspace, already open. If a mission has 5 core tools, "
         "the founder should see 5 tool cards when the mission loads."),
        ("Supporting",
         "Recommended. Sharpens or de-risks the answer. Show in a 'recommended tools' rail, "
         "one click away."),
        ("Stretch",
         "Optional. For a second pass or an ambitious founder. Show under 'go deeper', collapsed."),
        ("", ""),
        ("THE SHEETS", ""),
        ("Mission → Tools",
         "The main mapping. One row per mission-tool pair, with the reason and the specific "
         "mission question the tool's output should populate. This is the sheet to hand an engineer."),
        ("Coverage Grid",
         "15 columns × 122 rows. Read down a column to see a mission's whole toolset; read "
         "across a row to see a tool's whole lifecycle."),
        ("Tool → Missions",
         "Reverse index. For each tool: where it is used first, how many missions use it, and "
         "whether it is ever Core. Use this to decide which tools deserve the most polish."),
        ("Founder Toolkit",
         "The 21 tools with no natural home in the 15 missions, grouped into shelves with a "
         "trigger for when to surface each."),
        ("Resources",
         "Curated video / reading / template per mission, with a note on why each one earns its place."),
        ("Gaps & Recommendations",
         "What the mapping exposed about the mission design, and what to do about it."),
        ("", ""),
        ("SOURCING NOTE", ""),
        ("", "Every resource URL came from a live web search during compilation. None were "
             "written from memory. Direct page fetches were blocked by this environment's "
             "network policy, so links are search-index verified, not fetch-verified — run one "
             "link-check pass before publishing (scripts/check_links.py)."),
    ]
    r = 4
    for label, text in rows:
        if label and not text:
            c = ws.cell(row=r, column=2, value=label)
            c.font = Font(bold=True, size=11, color=ACCENT)
        elif label:
            c = ws.cell(row=r, column=2, value=label)
            c.font = Font(bold=True, size=10, color=INK)
            c.alignment = Alignment(vertical="top")
            c2 = ws.cell(row=r, column=3, value=text)
            c2.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = max(15, 13 * (len(text) // 95 + 1))
        elif text:
            c2 = ws.cell(row=r, column=3, value=text)
            c2.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = max(15, 13 * (len(text) // 95 + 1))
        r += 1
    return ws


# ---------------------------------------------------------------- Mission -> Tools
def sheet_mission_tools(wb):
    ws = wb.create_sheet("Mission → Tools")
    ws.sheet_view.showGridLines = False
    headers = ["Mission", "Mission title", "Phase", "Pts", "Role", "Tool",
               "Tool category", "Lives on tab", "Why this tool is here",
               "Feeds which mission question"]
    widths = [8, 26, 14, 6, 12, 32, 24, 22, 74, 40]
    header_row(ws, 1, headers, widths)

    r = 2
    for m in MISSIONS:
        block = MAPPING[m["id"]]
        for role_key in ("core", "supporting", "stretch"):
            for tool, why, feeds in block.get(role_key, []):
                _, cat, _, _ = TOOL_BY_NAME[tool]
                from startpad.tools import CATEGORY_TAB
                vals = [m["id"], m["title"], m["phase"], m["points"],
                        ROLE_LABEL[role_key], tool, cat, CATEGORY_TAB[cat], why, feeds]
                for i, v in enumerate(vals, start=1):
                    c = ws.cell(row=r, column=i, value=v)
                    c.border = BOX
                    c.alignment = Alignment(wrap_text=(i in (2, 6, 7, 8, 9, 10)),
                                            vertical="top", indent=1)
                    c.font = Font(size=10)
                ws.cell(row=r, column=5).fill = ROLE_FILL[ROLE_LABEL[role_key]]
                ws.cell(row=r, column=5).font = ROLE_FONT[ROLE_LABEL[role_key]]
                ws.cell(row=r, column=1).font = Font(size=10, bold=True)
                ws.cell(row=r, column=6).font = Font(size=10, bold=True, color=INK)
                ws.row_dimensions[r].height = max(15, 12 * (len(why) // 72 + 1))
                r += 1
    ws.auto_filter.ref = f"A1:J{r - 1}"
    return ws


# ---------------------------------------------------------------- Coverage grid
def sheet_grid(wb):
    ws = wb.create_sheet("Coverage Grid")
    ws.sheet_view.showGridLines = False

    lookup = {}
    for mid, block in MAPPING.items():
        for role_key in ("core", "supporting", "stretch"):
            for tool, _, _ in block.get(role_key, []):
                lookup[(tool, mid)] = ROLE_LABEL[role_key]

    ws.cell(row=1, column=1, value="Tool").font = Font(bold=True, color="FFFFFF", size=10)
    ws.cell(row=1, column=1).fill = PatternFill("solid", fgColor="2E6F5E")
    ws.cell(row=1, column=2, value="Category").font = Font(bold=True, color="FFFFFF", size=10)
    ws.cell(row=1, column=2).fill = PatternFill("solid", fgColor="2E6F5E")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 24
    for m in MISSIONS:
        col = 2 + m["id"]
        c = ws.cell(row=1, column=col, value=f'M{m["id"]}')
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor="2E6F5E")
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(col)].width = 6
        c2 = ws.cell(row=2, column=col, value=m["title"])
        c2.font = Font(size=8, color=INK)
        c2.alignment = Alignment(text_rotation=90, horizontal="center", vertical="bottom")
    ws.cell(row=2, column=1, value="").fill = SUB_FILL
    ws.row_dimensions[2].height = 118
    ws.freeze_panes = "C3"

    r = 3
    for cat in CATEGORY_ORDER:
        for name, tcat, _, _ in TOOLS:
            if tcat != cat:
                continue
            a = ws.cell(row=r, column=1, value=name)
            a.font = Font(size=10, bold=True)
            a.border = BOX
            a.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
            b = ws.cell(row=r, column=2, value=cat)
            b.font = Font(size=9, color="5A6B66")
            b.border = BOX
            b.alignment = Alignment(vertical="center", indent=1)
            hit = False
            for m in MISSIONS:
                col = 2 + m["id"]
                role = lookup.get((name, m["id"]))
                c = ws.cell(row=r, column=col)
                c.border = BOX
                c.alignment = Alignment(horizontal="center", vertical="center")
                if role:
                    hit = True
                    c.value = {"Core": "C", "Supporting": "S", "Stretch": "•"}[role]
                    c.fill = ROLE_FILL[role]
                    c.font = ROLE_FONT[role]
            if not hit:
                a.font = Font(size=10, bold=True, color="9AA6A2")
            r += 1

    r += 1
    ws.cell(row=r, column=1, value="LEGEND").font = Font(bold=True, size=10, color=ACCENT)
    r += 1
    for k, txt in (("C", "Core — required for the mission"),
                   ("S", "Supporting — recommended"),
                   ("•", "Stretch — optional / go deeper")):
        role = {"C": "Core", "S": "Supporting", "•": "Stretch"}[k]
        c = ws.cell(row=r, column=1, value=k)
        c.fill = ROLE_FILL[role]
        c.font = ROLE_FONT[role]
        c.alignment = Alignment(horizontal="center")
        c.border = BOX
        ws.cell(row=r, column=2, value=txt).font = Font(size=10)
        r += 1
    ws.cell(row=r, column=1, value="grey")
    ws.cell(row=r, column=1).font = Font(size=10, color="9AA6A2", bold=True)
    ws.cell(row=r, column=2,
            value="Tool name in grey = not used in any mission; lives on a Founder Toolkit shelf"
            ).font = Font(size=10)
    return ws


# ---------------------------------------------------------------- Tool -> Missions
def sheet_tool_missions(wb):
    ws = wb.create_sheet("Tool → Missions")
    ws.sheet_view.showGridLines = False
    headers = ["Tool", "Category", "What it does", "Framework reference",
               "First used", "# missions", "Core in", "Supporting in", "Stretch in",
               "Priority"]
    widths = [32, 24, 42, 40, 11, 11, 18, 22, 18, 12]
    header_row(ws, 1, headers, widths)

    per_tool = {}
    for mid, block in MAPPING.items():
        for role_key in ("core", "supporting", "stretch"):
            for tool, _, _ in block.get(role_key, []):
                per_tool.setdefault(tool, {"core": [], "supporting": [], "stretch": []})
                per_tool[tool][role_key].append(mid)

    r = 2
    for cat in CATEGORY_ORDER:
        for name, tcat, does, ref in TOOLS:
            if tcat != cat:
                continue
            d = per_tool.get(name, {"core": [], "supporting": [], "stretch": []})
            allm = sorted(set(d["core"] + d["supporting"] + d["stretch"]))
            first = f"M{min(allm)}" if allm else "—"
            n = len(allm)
            if d["core"]:
                prio = "P1 — build first" if len(d["core"]) > 1 else "P1"
            elif d["supporting"]:
                prio = "P2"
            elif d["stretch"]:
                prio = "P3"
            else:
                prio = "P4 — shelf"
            vals = [name, cat, does, ref, first, n,
                    ", ".join(f"M{i}" for i in sorted(d["core"])) or "—",
                    ", ".join(f"M{i}" for i in sorted(d["supporting"])) or "—",
                    ", ".join(f"M{i}" for i in sorted(d["stretch"])) or "—",
                    prio]
            for i, v in enumerate(vals, start=1):
                c = ws.cell(row=r, column=i, value=v)
                c.border = BOX
                c.font = Font(size=10)
                c.alignment = Alignment(wrap_text=(i in (1, 3, 4, 7, 8, 9)),
                                        vertical="top", indent=1)
            ws.cell(row=r, column=1).font = Font(size=10, bold=True)
            ws.cell(row=r, column=6).alignment = Alignment(horizontal="center")
            pc = ws.cell(row=r, column=10)
            if prio.startswith("P1"):
                pc.fill = ROLE_FILL["Core"]
                pc.font = ROLE_FONT["Core"]
            elif prio == "P2":
                pc.fill = ROLE_FILL["Supporting"]
                pc.font = ROLE_FONT["Supporting"]
            elif prio == "P3":
                pc.fill = ROLE_FILL["Stretch"]
                pc.font = ROLE_FONT["Stretch"]
            else:
                pc.fill = PatternFill("solid", fgColor="EFEFEF")
                pc.font = Font(size=10, color="6B7280", bold=True)
            r += 1
    ws.auto_filter.ref = f"A1:J{r - 1}"
    return ws


# ---------------------------------------------------------------- Founder toolkit
def sheet_toolkit(wb):
    ws = wb.create_sheet("Founder Toolkit")
    ws.sheet_view.showGridLines = False
    title_block(ws, "Founder Toolkit — tools that live outside the 15 missions",
                "21 tools have no natural home in the mission flow. Shelve them; don't force them in.", 4)
    headers = ["Shelf / when to surface it", "Tool", "Category", "Why it's here"]
    widths = [46, 34, 26, 68]
    header_row(ws, 4, headers, widths)
    ws.freeze_panes = "A5"

    r = 5
    for shelf, items in ALWAYS_AVAILABLE.items():
        start = r
        for tool, why in items:
            _, cat, _, _ = TOOL_BY_NAME[tool]
            for i, v in enumerate([shelf if r == start else "", tool, cat, why], start=1):
                c = ws.cell(row=r, column=i, value=v)
                c.border = BOX
                c.font = Font(size=10, bold=(i == 2))
                c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
            ws.cell(row=r, column=1).font = Font(size=10, bold=True, color=ACCENT)
            r += 1
        if r > start:
            ws.merge_cells(start_row=start, start_column=1, end_row=r - 1, end_column=1)
            ws.cell(row=start, column=1).alignment = Alignment(
                wrap_text=True, vertical="top", indent=1)
        r += 1
    return ws


# ---------------------------------------------------------------- Resources
def sheet_resources(wb):
    ws = wb.create_sheet("Resources")
    ws.sheet_view.showGridLines = False
    headers = ["Mission", "Mission title", "Type", "Title", "Source", "Length",
               "Lang", "Link", "Why this one"]
    widths = [10, 26, 11, 52, 24, 15, 7, 62, 66]
    header_row(ws, 1, headers, widths)

    r = 2

    def emit(mlabel, mtitle, items):
        nonlocal r
        for kind, title, source, url, length, lang, why in items:
            vals = [mlabel, mtitle, kind, title, source, length, lang, url, why]
            for i, v in enumerate(vals, start=1):
                c = ws.cell(row=r, column=i, value=v)
                c.border = BOX
                c.font = Font(size=10)
                c.alignment = Alignment(wrap_text=(i in (2, 4, 8, 9)), vertical="top", indent=1)
            ws.cell(row=r, column=1).font = Font(size=10, bold=True)
            ws.cell(row=r, column=4).font = Font(size=10, bold=True)
            lk = ws.cell(row=r, column=8)
            lk.hyperlink = url
            lk.font = Font(size=9, color="1D6FA3", underline="single")
            kc = ws.cell(row=r, column=3)
            if kind in ("Video", "Playlist"):
                kc.fill = PatternFill("solid", fgColor="FBE3E3")
                kc.font = Font(size=10, bold=True, color="A03434")
            elif kind in ("Course", "Guide"):
                kc.fill = PatternFill("solid", fgColor="E4EEF8")
                kc.font = Font(size=10, bold=True, color=SUPP)
            else:
                kc.fill = PatternFill("solid", fgColor="EDEDE6")
                kc.font = Font(size=10, bold=True, color=STRETCH)
            ws.row_dimensions[r].height = max(15, 12 * (len(why) // 64 + 1))
            r += 1

    emit("ALL", "Foundation — put on /resources", FOUNDATION)
    for m in MISSIONS:
        emit(f'M{m["id"]}', m["title"], BY_MISSION.get(m["id"], []))
    emit("M16*", "Raise Your Round (proposed)", FUNDRAISING)
    ws.auto_filter.ref = f"A1:I{r - 1}"
    return ws


# ---------------------------------------------------------------- Gaps
def sheet_gaps(wb, findings):
    ws = wb.create_sheet("Gaps & Recommendations")
    ws.sheet_view.showGridLines = False
    title_block(ws, "What the mapping exposed",
                "Findings from placing all 122 tools against the 15 missions.", 4)
    headers = ["#", "Finding", "Why it matters", "Recommendation"]
    widths = [5, 46, 62, 62]
    header_row(ws, 4, headers, widths)
    ws.freeze_panes = "A5"
    r = 5
    for i, (finding, why, rec) in enumerate(findings, start=1):
        for j, v in enumerate([i, finding, why, rec], start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.border = BOX
            c.font = Font(size=10, bold=(j == 2))
            c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
        ws.row_dimensions[r].height = max(30, 12 * (max(len(why), len(rec)) // 58 + 1))
        r += 1
    return ws


FINDINGS = [
    ("The library holds 122 unique tools, not 123",
     "'Term Sheet Analyzer' is listed on both the Cap Table & Deal Terms tab and the "
     "Fundraising & Investor tab. Counting the listing twice inflates the headline number, "
     "and in a database it would become two rows a founder can fill in inconsistently.",
     "Keep one canonical tool record and cross-link it from both tabs. Publish '122 tools' "
     "or '123 tool placements' — pick one and use it consistently in marketing copy."),
    ("21 tools have no home in the 15 missions",
     "Almost the entire fundraising block (13 tools), most of the finance statements, and the "
     "team/governance tools are never reachable through the mission flow. A founder who only "
     "follows missions will never see the Cap Table, the Pitch Deck template, or the Legal "
     "Checklist — even though those are among the highest-value assets in the library.",
     "Ship a 'Founder Toolkit' shelf that is always reachable from the dashboard, and add a "
     "proposed Mission 16 — 'Raise Your Round' — that gives the 15 fundraising tools a "
     "narrative home. See the Founder Toolkit sheet."),
    ("Missions 7 and 10 have the same title",
     "Both are called 'Solution Hypothesis Test'. In a sidebar or a progress list they are "
     "indistinguishable, and the AI evaluator has no way to tell a founder that the second one "
     "expects prototype evidence rather than conversation evidence.",
     "Rename Mission 10 to 'Prototype Hypothesis Test' (or 'Solution Test — with Prototype'). "
     "The mapping in this workbook already treats them as distinct: M7 is interview- and "
     "smoke-test-led, M10 is usability-test-led."),
    ("Mission 4 Q4 asks how you'll measure success, before any metric exists",
     "Founders answer this with a vague sentence because no tool is attached. It is the single "
     "cheapest place to introduce the North Star Metric — and doing it here means Missions 13 "
     "and 14 have something to instrument against.",
     "Attach North Star Metric and KPI Scorecard as Core in Mission 4, and carry the founder's "
     "answer forward as a pre-filled value in Missions 13 and 14."),
    ("Mission 13 ships an MVP with no analytics requirement",
     "Mission 14 then asks for retention rate, user metrics and a PMF score. If nothing was "
     "instrumented during the build, Mission 14 is unanswerable and the founder either stalls "
     "or invents numbers — which the AI evaluator will score as low-evidence.",
     "Make AARRR Funnel and North Star Metric a gating checklist item in Mission 13 Q6 "
     "('Is your MVP ready for launch?'). No event taxonomy, no launch."),
    ("Mission 14 asks the founder to self-rate product-market fit 1-10",
     "A self-rating is exactly the kind of unfalsifiable answer the platform's 5-pillar rubric "
     "(Specificity, Depth, Evidence, Actionability, Relevance) is designed to catch. The founder "
     "is being set up to score badly on Evidence.",
     "Replace or supplement Q5 with the Sean Ellis 40% survey ('How would you feel if you could "
     "no longer use this product?'). It produces a real, benchmarked number and the resource "
     "sheet links the survey template and the Superhuman method article."),
    ("Mission 11 ('Service Selection') has a title that doesn't describe it",
     "The mission body is a go-to-market strategy — target market, channels, pricing, sales "
     "process, launch plan. 'Service Selection' reads like a vendor-picking exercise and will "
     "cost completion rate at the point a founder decides whether to start it.",
     "Rename to 'Go-to-Market Strategy' and rename Mission 15 to 'Launch & Growth' so the two "
     "are clearly plan-then-execute rather than two go-to-market missions."),
    ("Mission 15 is heavily overloaded",
     "24 tools map to it — more than any other mission, and roughly a fifth of the library. "
     "At 500 points and 10-20 hours it is already the largest mission; presenting 24 tool cards "
     "will read as a wall rather than a path.",
     "Split the surfaced set: Core (7) on the mission screen, Supporting behind a 'recommended' "
     "rail, and move the standing metrics tools (SaaS Dashboard, KPI Scorecard, Rule of 40, "
     "Magic Number) to the post-launch Scale shelf where they belong."),
    ("The strongest tools are back-loaded; the first four missions are thin",
     "Missions 1-4 together carry 500 of the 4,050 points and only ~12% of tool usage, but they "
     "are where founders drop off. Discovery is also the phase where a bad answer is cheapest "
     "to fix and most expensive to leave uncorrected.",
     "Weight the early missions with more guided tooling (the mapping gives M1 five Core tools) "
     "and put the highest-quality video assets in the pre-mission briefing for M1 and M2 — "
     "the YC 'How to talk to users' and JTBD milkshake resources are the two highest-leverage "
     "pieces in the whole library."),
    ("Nine tools are Core in three or more missions",
     "Problem Statement, Assumptions Tracker, Customer Persona, Interview Script Builder, "
     "Validation Scorecard, MVP Scope, North Star Metric, Hypothesis Builder and "
     "Interview Synthesis recur across the journey. Today each mission would present them as a "
     "blank tool.",
     "Make these tools stateful across missions: a founder's Mission 2 persona should appear "
     "pre-filled in Missions 4, 11 and 12, editable, with a visible version history. This is the "
     "single biggest UX win available from this mapping — see the 'First used' and "
     "'# missions' columns on the Tool → Missions sheet."),
    ("Resource links are search-verified, not fetch-verified",
     "This environment's network policy blocked direct page fetches, so every URL here came "
     "from a live search index rather than from opening the page. Search indexes lag; a video "
     "can be made private after indexing.",
     "Run scripts/check_links.py from an unrestricted network before publishing any of these to "
     "startpad.me, and re-run it quarterly. Prefer the YC Library, Strategyzer, NN/g and HBR "
     "links as permanent anchors — they are the least likely to move."),
    ("There is no Arabic-language resource depth",
     "The platform is bilingual EN/AR with full RTL, but the high-quality startup education "
     "corpus is overwhelmingly English. Only three Arabic resources survived the quality bar "
     "used here, and one of them is a paid course.",
     "Treat Arabic resource creation as a build, not a curation, problem: commission or record "
     "short AR briefings for Missions 1-8 in-house. That is a genuine moat for a MENA-focused "
     "platform, and it is the thing competitors will not copy quickly."),
]


def build(out_path):
    wb = Workbook()
    wb.remove(wb.active)
    sheet_readme(wb)
    sheet_mission_tools(wb)
    sheet_grid(wb)
    sheet_tool_missions(wb)
    sheet_toolkit(wb)
    sheet_resources(wb)
    sheet_gaps(wb, FINDINGS)
    wb.properties.created = FIXED_TIMESTAMP
    wb.properties.modified = FIXED_TIMESTAMP
    wb.save(out_path)
    normalize_zip(out_path)
    return out_path


if __name__ == "__main__":
    p = build(sys.argv[1] if len(sys.argv) > 1 else "output/Startpad_Mission_Tool_Matrix.xlsx")
    print("wrote", p)
