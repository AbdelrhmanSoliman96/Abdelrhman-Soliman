# -*- coding: utf-8 -*-
"""
Render the StartPad template library to Word.

One .docx per template per language: cover, a page explaining the tool and where
it comes from, the fillable grid itself, and a closing page about StartPad.

BRAND RULES ENCODED HERE (from the 2026 brand presentation)
  Lime #CBF24A is a ground, never ink — at 1.29:1 on white it cannot carry type,
  so it fills band headings and rules and never sets a word of body copy.
  Ultraviolet #3418E0 carries white type at 8.9:1 and is therefore the surface
  for the cover. Ink #0B0E14 and white do the reading.
  Satoshi sets Latin, Zain sets Arabic. Word substitutes if they are not
  installed locally; the file still opens correctly.
  One language per asset — so English and Arabic are separate files, not
  facing columns.

Arabic needs three signals Word treats independently: w:bidi on the paragraph,
w:rtl on the run with a complex-script font and size, and w:bidiVisual on the
table so columns run right-to-left too. Section-level RTL is set as well, so
page furniture follows.
"""
import os
import sys

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from scripts import _ooxml as X  # noqa: E402
from scripts._determinism import FIXED_TIMESTAMP, normalize_zip  # noqa: E402

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LOGO = os.path.join(REPO, "brand", "logo")

LIME = "CBF24A"
UV = "3418E0"
INK_HEX = "0B0E14"
INK = RGBColor(0x0B, 0x0E, 0x14)
UV_RGB = RGBColor(0x34, 0x18, 0xE0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0x5E, 0x66, 0x70)
HAIR = "D8DCE2"

LATIN = "Satoshi"
ARABIC = "Zain"

TEXT_W = {"portrait": 17.0, "landscape": 25.7}


# ----------------------------------------------------------------- xml helpers
# Ordering matters: see scripts/_ooxml.py. Everything structural goes through
# those helpers so the files open in LibreOffice, Google Docs and Pages, not
# only in Word.
def style_run(run, size, rtl, bold=False, italic=False, color=INK, caps=False,
              spacing=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = LATIN
    if caps:
        run.font.all_caps = True
    if spacing is not None:
        X.run_spacing(run, spacing)
    X.run_cs(run, ARABIC if rtl else LATIN, size, rtl)
    return run


def para(doc, text="", size=10.5, rtl=False, bold=False, italic=False, color=INK,
         before=0, after=6, align=None, line=1.3, caps=False, indent=0,
         spacing=None, keep=False):
    p = doc.add_paragraph()
    if rtl:
        X.para_rtl(p)
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep
    if indent:
        if rtl:
            pf.right_indent = Cm(indent)
        else:
            pf.left_indent = Cm(indent)
    p.alignment = align if align is not None else (
        WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT)
    if text:
        style_run(p.add_run(text), size, rtl, bold, italic, color, caps, spacing)
    return p


def cell_para(cell, first=False, rtl=False, after=0, line=1.2):
    """A paragraph inside a cell, with RTL applied before spacing and jc."""
    p = cell.paragraphs[0] if first else cell.add_paragraph()
    if rtl:
        X.para_rtl(p)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if rtl:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    return p


def rule(doc, hexcolor=LIME, cm=0.22, after=10, before=0):
    t = doc.add_table(rows=1, cols=1)
    X.tbl_borders(t, style="none")
    c = t.rows[0].cells[0]
    X.tc_shade(c, hexcolor)
    X.tr_height(t.rows[0], cm)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    style_run(p.add_run(""), 1, False)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(after)
    spacer.paragraph_format.space_before = Pt(before)
    return t


# -------------------------------------------------------------------- section
def set_orientation(section, orient):
    if orient == "landscape":
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Cm(29.7), Cm(21.0)
        section.left_margin = section.right_margin = Cm(2.0)
        section.top_margin = Cm(1.6)
        section.bottom_margin = Cm(1.5)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
        section.left_margin = section.right_margin = Cm(2.0)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(1.7)


def footer_for(section, label, rtl):
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    if rtl:
        X.para_rtl(p)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_run(p.add_run(label), 7.5, rtl, color=GREY, spacing=12)


# ----------------------------------------------------------------------- pages
def cover(doc, t, lang, rtl, strings):
    band = doc.add_table(rows=1, cols=1)
    X.tbl_borders(band, style="none")
    c = band.rows[0].cells[0]
    X.tc_shade(c, UV)
    X.tc_margins(c, top=0.9, start=0.9, bottom=0.9, end=0.9)
    X.tr_height(band.rows[0], 11.4)
    c.paragraphs[0].paragraph_format.space_after = Pt(0)
    run = c.paragraphs[0].add_run()
    run.add_picture(os.path.join(LOGO, "startpad_horizontal_white.png"),
                    width=Cm(4.6))
    if rtl:
        X.para_rtl(c.paragraphs[0])
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    def band_p(text, size, bold=False, after=4, color=WHITE, before=0, caps=False,
               spacing=None, line=1.15):
        p = c.add_paragraph()
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.line_spacing = line
        if rtl:
            X.para_rtl(p)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        style_run(p.add_run(text), size, rtl, bold, color=color, caps=caps,
                  spacing=spacing)

    band_p(strings["label"], 8, bold=True, after=18, before=26, caps=not rtl,
           spacing=26 if not rtl else None)
    band_p(t["name"], 30 if not rtl else 27, bold=True, after=8, line=1.05)
    band_p(t["tagline"], 12.5, after=0, line=1.35)

    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    rule(doc, LIME, 0.3, after=12)
    para(doc, strings["produces"] + " " + t["produces"], 11, rtl, bold=True,
         after=4)
    para(doc, "startpad.me", 10, rtl, color=UV_RGB, bold=True, after=0)


def about_template(doc, t, rtl, s):
    para(doc, s["about_h"], 19, rtl, bold=True, after=2, keep=True)
    rule(doc, LIME, 0.18, after=10)

    para(doc, s["what_h"], 11.5, rtl, bold=True, color=UV_RGB, after=3,
         before=4, keep=True)
    para(doc, t["what"], 10.5, rtl, after=10)

    para(doc, s["purpose_h"], 11.5, rtl, bold=True, color=UV_RGB, after=3,
         before=4, keep=True)
    for b in t["purpose"]:
        para(doc, "—  " + b, 10, rtl, indent=0.5, after=3)

    para(doc, s["missions_h"], 11.5, rtl, bold=True, color=UV_RGB, after=3,
         before=10, keep=True)
    para(doc, t["missions"], 10.5, rtl, after=10)

    para(doc, s["ref_h"], 11.5, rtl, bold=True, color=UV_RGB, after=3,
         before=4, keep=True)
    para(doc, s["ref_primary"], 9, rtl, bold=True, color=GREY, after=1)
    para(doc, t["ref_primary"], 10, rtl, after=6)
    para(doc, s["ref_origin"], 9, rtl, bold=True, color=GREY, after=1)
    para(doc, t["ref_origin"], 10, rtl, after=6)
    para(doc, s["ref_further"], 9, rtl, bold=True, color=GREY, after=1)
    for r in t["ref_further"]:
        para(doc, "—  " + r, 9.5, rtl, indent=0.5, after=2)

    para(doc, s["how_h"], 11.5, rtl, bold=True, color=UV_RGB, after=3,
         before=12, keep=True)
    for i, step in enumerate(t["how"], 1):
        para(doc, f"{i}.  " + step, 10, rtl, indent=0.5, after=3)

    para(doc, s["done_h"], 11.5, rtl, bold=True, color=UV_RGB, after=3,
         before=12, keep=True)
    for b in t["done_well"]:
        para(doc, "☐  " + b, 10, rtl, indent=0.5, after=3)


def grid(doc, cols, spec, rtl):
    rows = len(spec)
    table = doc.add_table(rows=rows, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    X.tbl_borders(table)
    if rtl:
        X.tbl_rtl(table)
    occupied = [[False] * cols for _ in range(rows)]

    for r, spec_row in enumerate(spec):
        tallest = 0
        for box in spec_row:
            w, h = box.get("w", 1), box.get("h", 1)
            c = next((i for i in range(cols) if not occupied[r][i]), None)
            if c is None or c + w > cols:
                continue
            for rr in range(r, min(r + h, rows)):
                for cc in range(c, c + w):
                    occupied[rr][cc] = True
            target = table.cell(min(r + h - 1, rows - 1), c + w - 1)
            cell = table.cell(r, c).merge(target) if (w > 1 or h > 1) \
                else table.cell(r, c)
            X.tc_margins(cell)
            head = box.get("head")
            if head:
                X.tc_shade(cell, LIME)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2 if box.get("q") else 0)
            if rtl:
                X.para_rtl(p)
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            style_run(p.add_run(box["t"]), 9.5 if not head else 10, rtl,
                      bold=True, color=INK, spacing=6 if head else None)
            if box.get("q"):
                q = cell.add_paragraph()
                q.paragraph_format.space_after = Pt(0)
                q.paragraph_format.line_spacing = 1.15
                if rtl:
                    X.para_rtl(q)
                    q.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                style_run(q.add_run(box["q"]), 7.8, rtl, italic=True, color=GREY)
            tallest = max(tallest, box.get("cm", 3.0) / max(h, 1))
        X.tr_height(table.rows[r], tallest or 1.0)
    return table


def log_table(doc, spec, rtl, s):
    para(doc, spec["title"], 12, rtl, bold=True, color=UV_RGB, before=14,
         after=2, keep=True)
    if spec.get("note"):
        para(doc, spec["note"], 8.5, rtl, italic=True, color=GREY, after=6)
    seed = spec.get("seed") or []
    n = spec.get("rows", 0) + len(seed)
    table = doc.add_table(rows=n + 1, cols=len(spec["cols"]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    X.tbl_borders(table)
    if rtl:
        X.tbl_rtl(table)
    for i, (head, w) in enumerate(zip(spec["cols"], spec["widths"])):
        cell = table.rows[0].cells[i]
        X.tc_shade(cell, LIME)
        X.tc_margins(cell, top=0.1, bottom=0.1)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        if rtl:
            X.para_rtl(p)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        style_run(p.add_run(head), 8, rtl, bold=True, color=INK)
    X.tr_height(table.rows[0], 0.9)
    for r in range(1, n + 1):
        X.tr_height(table.rows[r], 1.05)
        for i in range(len(spec["cols"])):
            cell = table.rows[r].cells[i]
            X.tc_margins(cell, top=0.1, bottom=0.1)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if rtl:
                X.para_rtl(p)
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if r - 1 < len(seed):
                style_run(p.add_run(seed[r - 1][i]), 8.5, rtl, color=INK)
            elif i == 0 and spec["cols"][0] == "#":
                style_run(p.add_run(str(r)), 8.5, rtl, color=GREY)
    for i, w in enumerate(spec["widths"]):
        for row in table.rows:
            row.cells[i].width = Cm(w)
    return table


def about_startpad(doc, rtl, s):
    band = doc.add_table(rows=1, cols=1)
    X.tbl_borders(band, style="none")
    c = band.rows[0].cells[0]
    X.tc_shade(c, LIME)
    X.tc_margins(c, top=0.7, start=0.7, bottom=0.7, end=0.7)
    X.tr_height(band.rows[0], 4.2)
    p0 = c.paragraphs[0]
    p0.paragraph_format.space_after = Pt(6)
    if rtl:
        X.para_rtl(p0)
        p0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p0.add_run().add_picture(os.path.join(LOGO, "startpad_horizontal_ink.png"),
                             width=Cm(4.4))
    p1 = c.add_paragraph()
    p1.paragraph_format.space_after = Pt(0)
    if rtl:
        X.para_rtl(p1)
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_run(p1.add_run(s["sp_line"]), 12.5, rtl, bold=True, color=INK)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    for block in s["sp_body"]:
        para(doc, block, 10.5, rtl, after=8)
    para(doc, s["sp_missions_h"], 11, rtl, bold=True, color=UV_RGB, before=6,
         after=3)
    for b in s["sp_missions"]:
        para(doc, "—  " + b, 10, rtl, indent=0.5, after=3)
    rule(doc, LIME, 0.18, after=8, before=8)
    para(doc, "startpad.me", 13, rtl, bold=True, color=UV_RGB, after=2)
    para(doc, s["sp_foot"], 9, rtl, color=GREY, after=0)


# ------------------------------------------------------------------ assembling
def needs_landscape(t):
    """Wide grids and wide logs do not fit A4 portrait's 17cm text column."""
    if t.get("cols", 1) >= 5 or t.get("cols2", 1) >= 5:
        return True
    log = t.get("log")
    if log and sum(log["widths"]) > TEXT_W["portrait"] - 0.2:
        return True
    return False


def fit_widths(table, total_cm, cols):
    """Word ignores cell widths unless every cell in the column carries one."""
    w = Cm(total_cm / cols)
    for row in table.rows:
        for cell in row.cells:
            cell.width = w


def build_one(t, lang, strings, out_path):
    rtl = lang == "ar"
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = LATIN
    st.font.size = Pt(10.5)
    st.font.color.rgb = INK

    sec = doc.sections[0]
    set_orientation(sec, "portrait")
    if rtl:
        X.sect_rtl(sec)
    label = "%s · %s · startpad.me" % (strings["brand"], t["name"])
    footer_for(sec, label, rtl)

    cover(doc, t, lang, rtl, strings)
    doc.add_page_break()
    about_template(doc, t, rtl, strings)

    land = needs_landscape(t)
    sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
    set_orientation(sec2, "landscape" if land else "portrait")
    if rtl:
        X.sect_rtl(sec2)
    footer_for(sec2, label, rtl)
    width = TEXT_W["landscape" if land else "portrait"]

    para(doc, strings["fill_h"], 15, rtl, bold=True, after=2, keep=True)
    para(doc, strings["fill_note"], 9, rtl, italic=True, color=GREY, after=8)
    g = grid(doc, t["cols"], t["grid"], rtl)
    fit_widths(g, width, t["cols"])

    if t.get("log"):
        lg = log_table(doc, t["log"], rtl, strings)
        for i, w in enumerate(t["log"]["widths"]):
            for row in lg.rows:
                row.cells[i].width = Cm(w * width / sum(t["log"]["widths"]))

    if t.get("grid2"):
        para(doc, t.get("grid2_title", ""), 12, rtl, bold=True, color=UV_RGB,
             before=16, after=4, keep=True)
        g2 = grid(doc, t["cols2"], t["grid2"], rtl)
        fit_widths(g2, width, t["cols2"])

    sec3 = doc.add_section(WD_SECTION.NEW_PAGE)
    set_orientation(sec3, "portrait")
    if rtl:
        X.sect_rtl(sec3)
    footer_for(sec3, label, rtl)
    about_startpad(doc, rtl, strings)

    core = doc.core_properties
    core.title = "%s — StartPad template" % t["name"]
    core.author = "StartPad"
    core.category = "StartPad founder template"
    core.created = FIXED_TIMESTAMP
    core.modified = FIXED_TIMESTAMP

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    doc.save(out_path)
    normalize_zip(out_path)
    return out_path


def build_all(out_root):
    from templates.catalog_en import EN
    from templates.strings import STRINGS
    written = []
    order = list(EN)
    try:
        from templates.catalog_ar import AR
    except ImportError:
        AR = {}
    for lang, catalog in (("en", EN), ("ar", AR)):
        if not catalog:
            continue
        for i, slug in enumerate(order, 1):
            if slug not in catalog:
                continue
            out = os.path.join(out_root, lang,
                               "%02d_%s_%s.docx" % (i, slug, lang))
            build_one(catalog[slug], lang, STRINGS[lang], out)
            written.append(out)
    return written


if __name__ == "__main__":
    root = sys.argv[1] if len(sys.argv) > 1 else "output/templates"
    for f in build_all(root):
        print("wrote", f, os.path.getsize(f), "bytes")
