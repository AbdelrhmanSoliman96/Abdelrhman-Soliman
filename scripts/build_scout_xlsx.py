"""Build the MENA startup-programme scouting registry workbook."""
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from scout.sources import (  # noqa: E402
    COVERAGE_NOTES, ENTITY_TYPES, PROGRAM_TYPES, PROGRAMS, STAGES,
)
from scripts._determinism import FIXED_TIMESTAMP, normalize_zip  # noqa: E402

INK = "16211E"
ACCENT = "0F5C4B"
HEAD = PatternFill("solid", fgColor="0F5C4B")
SUBHEAD = PatternFill("solid", fgColor="2E6F5E")
SUNK = PatternFill("solid", fgColor="ECF0EE")
THIN = Side(style="thin", color="D9E2DE")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CONF_FILL = {
    "high": PatternFill("solid", fgColor="D7EBE3"),
    "medium": PatternFill("solid", fgColor="F5EAD5"),
    "low": PatternFill("solid", fgColor="F7DDD3"),
}
CONF_FONT = {
    "high": Font(size=10, bold=True, color="0F5C4B"),
    "medium": Font(size=10, bold=True, color="8A6E3C"),
    "low": Font(size=10, bold=True, color="B4531C"),
}


def slug(s):
    s = re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")
    return re.sub(r"-+", "-", s)[:60]


def banner(ws, title, subtitle, width):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    c = ws.cell(row=1, column=1, value=title)
    c.font = Font(size=15, bold=True, color="FFFFFF")
    c.fill = HEAD
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 28
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    c = ws.cell(row=2, column=1, value=subtitle)
    c.font = Font(size=10, italic=True, color=INK)
    c.fill = SUNK
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[2].height = 19


def header(ws, row, headers, widths, freeze_col=1):
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = SUBHEAD
        c.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
        c.border = BOX
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row=row + 1, column=freeze_col)


# ------------------------------------------------------------------ README
def sheet_readme(wb):
    ws = wb.create_sheet("READ ME FIRST")
    ws.sheet_view.showGridLines = False
    banner(ws, "MENA Startup Programme Scout — source registry",
           "Egypt · GCC · wider MENA.  Built to be read by a scraper, not just by a person.", 5)
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 108

    blocks = [
        ("WHAT THIS IS", ""),
        ("", "A registry of the organisations that run startup programmes across Egypt, the "
             "GCC and the wider MENA region, and of the programmes themselves. One row per "
             "programme — entities repeat by design, because most run more than one."),
        ("", ""),
        ("HOW THE AUTOMATION WORKS", ""),
        ("", "The SCRAPE_CONFIG sheet is the machine-readable half of this file. "
             "scripts/scout_scraper.py reads it, visits each source, and writes back what it "
             "finds. You do not hand-write CSS selectors: the scraper auto-discovers content "
             "in priority order — RSS/Atom feed, then sitemap.xml, then JSON-LD structured "
             "data, then a heading/link heuristic on the page."),
        ("", "It stores a fingerprint of every item it has seen. On the next run, anything "
             "whose fingerprint is new is written to output/scout/new_programs.csv. That is "
             "the answer to 'if there is any new programme, scrape it automatically'."),
        ("", ""),
        ("THE SHEETS", ""),
        ("PROGRAMS", "The main table. Entity, programme, type, stage, funding, equity, "
                     "duration, description, stats, URLs, and a confidence rating. This is "
                     "the sheet to read as a human."),
        ("ENTITIES", "Deduplicated organisations, with a count of how many programmes each "
                     "runs. Use this to prioritise outreach."),
        ("SCRAPE_CONFIG", "What the scraper reads. One row per URL to watch, with method, "
                          "frequency and an enabled flag. Add a row here and the scraper "
                          "picks it up on the next run — no code change."),
        ("VOCAB", "Controlled vocabularies. The PROGRAMS sheet has dropdown validation "
                  "wired to these, so hand-added rows stay machine-readable."),
        ("COVERAGE & GAPS", "What this file does and does not cover, stated plainly."),
        ("", ""),
        ("READ THE CONFIDENCE COLUMN", ""),
        ("high", "Entity and URL both appeared as a live search-result link, and the facts "
                 "come from the result text."),
        ("medium", "Entity confirmed in result text; URL inferred from another page on the "
                   "same domain."),
        ("low", "Entity named in prose only — the URL is a plausible guess. Verify these first."),
        ("", ""),
        ("IMPORTANT — NOTHING HERE IS FETCH-VERIFIED", ""),
        ("", "The environment this was compiled in blocked every outbound request. Direct "
             "connection attempts to flat6labs.com, magnitt.com, wamda.com, itida.gov.eg, "
             "hub71.com, oasis500.com, sheraa.ae and startupqatar.qa all failed. Every URL "
             "came from a live web search, but no page was opened. Run the scraper from a "
             "normal network before treating any row as confirmed — it writes an http_status "
             "back for every source."),
        ("", ""),
        ("FIRST RUN", ""),
        ("", "pip install requests beautifulsoup4 lxml openpyxl"),
        ("", "python3 scripts/scout_scraper.py --verify-only     # just check every URL resolves"),
        ("", "python3 scripts/scout_scraper.py                   # full discovery pass"),
        ("", "python3 scripts/scout_scraper.py --since-last      # only report what is new"),
    ]
    r = 4
    for label, text in blocks:
        if label and not text:
            c = ws.cell(row=r, column=2, value=label)
            c.font = Font(bold=True, size=11, color=ACCENT)
        elif label:
            c = ws.cell(row=r, column=2, value=label)
            c.font = Font(bold=True, size=10, color=INK)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c2 = ws.cell(row=r, column=3, value=text)
            c2.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = max(15, 13 * (len(text) // 105 + 1))
        elif text:
            c2 = ws.cell(row=r, column=3, value=text)
            c2.alignment = Alignment(wrap_text=True, vertical="top")
            if text.startswith(("pip ", "python3 ")):
                c2.font = Font(name="Consolas", size=9.5, color="1D6FA3")
            ws.row_dimensions[r].height = max(15, 13 * (len(text) // 105 + 1))
        r += 1
    return ws


# ------------------------------------------------------------------ PROGRAMS
COLS = [
    ("entity", "Entity", 30), ("entity_type", "Entity type", 22),
    ("country", "Country", 16), ("city", "City", 16),
    ("program", "Programme", 34), ("program_type", "Programme type", 20),
    ("stage", "Stage", 12), ("sectors", "Sectors", 26),
    ("eligibility", "Eligibility", 26), ("funding", "Funding", 30),
    ("equity", "Equity", 16), ("duration", "Duration", 16),
    ("cadence", "Cadence", 20), ("status", "Status", 20),
    ("description", "Description", 66), ("stats", "Stats / proof points", 40),
    ("entity_url", "Entity URL", 40), ("programs_url", "Programmes URL (scrape me)", 44),
    ("apply_url", "Apply URL", 36), ("confidence", "Confidence", 12),
    ("source", "Source / provenance", 46),
]


def sheet_programs(wb):
    ws = wb.create_sheet("PROGRAMS")
    ws.sheet_view.showGridLines = False
    header(ws, 1, [c[1] for c in COLS], [c[2] for c in COLS], freeze_col=3)

    rows = sorted(PROGRAMS, key=lambda p: (p["country"], p["entity"], p["program"]))
    r = 2
    for p in rows:
        for i, (key, _, _) in enumerate(COLS, start=1):
            v = p[key]
            c = ws.cell(row=r, column=i, value=v)
            c.border = BOX
            c.font = Font(size=10)
            c.alignment = Alignment(wrap_text=key in {
                "entity", "program", "sectors", "eligibility", "funding", "description",
                "stats", "entity_url", "programs_url", "apply_url", "source", "cadence",
                "status"}, vertical="top", indent=1)
            if key in {"entity_url", "programs_url", "apply_url"} and v:
                c.hyperlink = v
                c.font = Font(size=9, color="1D6FA3", underline="single")
        ws.cell(row=r, column=1).font = Font(size=10, bold=True)
        ws.cell(row=r, column=5).font = Font(size=10, bold=True, color=ACCENT)
        conf = ws.cell(row=r, column=20)
        conf.fill = CONF_FILL[p["confidence"]]
        conf.font = CONF_FONT[p["confidence"]]
        conf.alignment = Alignment(horizontal="center", vertical="top")
        if str(p["status"]).startswith("CLOSED"):
            ws.cell(row=r, column=14).fill = PatternFill("solid", fgColor="F0D5CC")
            ws.cell(row=r, column=14).font = Font(size=10, bold=True, color="B4531C")
        ws.row_dimensions[r].height = max(28, 11.5 * (len(p["description"]) // 62 + 1))
        r += 1

    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{r - 1}"

    # dropdown validation so hand-added rows stay machine-readable
    last = r + 400
    for col_idx, vocab in ((2, ENTITY_TYPES), (6, PROGRAM_TYPES), (7, STAGES),
                           (20, ["high", "medium", "low"])):
        dv = DataValidation(type="list", formula1='"' + ",".join(vocab) + '"',
                            allow_blank=True, showDropDown=False)
        dv.error = "Use a value from the VOCAB sheet so the scraper can parse this column."
        dv.errorTitle = "Controlled vocabulary"
        ws.add_data_validation(dv)
        col = get_column_letter(col_idx)
        dv.add(f"{col}2:{col}{last}")
    return ws, r - 1


# ------------------------------------------------------------------ ENTITIES
def sheet_entities(wb):
    ws = wb.create_sheet("ENTITIES")
    ws.sheet_view.showGridLines = False
    heads = ["Entity", "Entity type", "Country", "City", "# programmes",
             "Programmes", "Entity URL", "Best confidence"]
    header(ws, 1, heads, [32, 22, 18, 18, 12, 52, 42, 14])

    agg = {}
    for p in PROGRAMS:
        k = (p["entity"], p["country"])
        e = agg.setdefault(k, {"type": p["entity_type"], "city": p["city"],
                               "progs": [], "url": p["entity_url"], "conf": []})
        e["progs"].append(p["program"])
        e["conf"].append(p["confidence"])
        if not e["url"] and p["entity_url"]:
            e["url"] = p["entity_url"]

    order = {"high": 0, "medium": 1, "low": 2}
    r = 2
    for (name, country), e in sorted(agg.items(), key=lambda kv: (kv[0][1], kv[0][0])):
        best = sorted(e["conf"], key=lambda c: order[c])[0]
        vals = [name, e["type"], country, e["city"], len(e["progs"]),
                " · ".join(e["progs"]), e["url"], best]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = BOX
            c.font = Font(size=10)
            c.alignment = Alignment(wrap_text=i in (1, 6, 7), vertical="top", indent=1)
        ws.cell(row=r, column=1).font = Font(size=10, bold=True)
        ws.cell(row=r, column=5).alignment = Alignment(horizontal="center")
        if e["url"]:
            lk = ws.cell(row=r, column=7)
            lk.hyperlink = e["url"]
            lk.font = Font(size=9, color="1D6FA3", underline="single")
        bc = ws.cell(row=r, column=8)
        bc.fill = CONF_FILL[best]
        bc.font = CONF_FONT[best]
        bc.alignment = Alignment(horizontal="center")
        r += 1
    ws.auto_filter.ref = f"A1:H{r - 1}"
    return ws, r - 1


# ------------------------------------------------------------------ SCRAPE_CONFIG
def sheet_scrape_config(wb):
    ws = wb.create_sheet("SCRAPE_CONFIG")
    ws.sheet_view.showGridLines = False
    banner(ws, "SCRAPE_CONFIG — the machine-readable half of this workbook",
           "scripts/scout_scraper.py reads this sheet. Add a row and it is picked up on the "
           "next run — no code change needed.", 12)
    heads = ["source_id", "entity", "country", "url", "url_role", "method",
             "frequency_days", "enabled", "priority", "expect", "last_run",
             "last_http_status", "notes"]
    widths = [30, 28, 16, 50, 16, 14, 14, 10, 10, 30, 14, 16, 44]
    header(ws, 4, heads, widths, freeze_col=2)

    seen, rows = set(), []
    for p in sorted(PROGRAMS, key=lambda x: (x["country"], x["entity"])):
        for url, role in ((p["programs_url"], "programs"),
                          (p["apply_url"], "apply"),
                          (p["entity_url"], "entity")):
            if not url or url in seen:
                continue
            seen.add(url)
            is_feed = p["program_type"] in ("News feed", "Directory listing")
            prio = 1 if is_feed else (2 if role == "programs" else 3)
            rows.append([
                f"{slug(p['entity'])}--{role}", p["entity"], p["country"], url, role,
                "auto", 7 if is_feed else 14, "yes", prio,
                "programme calls / cohort news" if is_feed else "programme detail",
                "", "", "auto = try RSS, then sitemap, then JSON-LD, then heading heuristic",
            ])

    rows.sort(key=lambda x: (x[8], x[2], x[1]))
    r = 5
    for row in rows:
        for i, v in enumerate(row, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = BOX
            c.font = Font(size=9.5)
            c.alignment = Alignment(wrap_text=i in (1, 2, 4, 10, 13), vertical="top", indent=1)
        ws.cell(row=r, column=1).font = Font(name="Consolas", size=9)
        lk = ws.cell(row=r, column=4)
        lk.hyperlink = row[3]
        lk.font = Font(size=9, color="1D6FA3", underline="single")
        pc = ws.cell(row=r, column=9)
        pc.alignment = Alignment(horizontal="center")
        pc.fill = {1: CONF_FILL["high"], 2: CONF_FILL["medium"], 3: SUNK}[row[8]]
        ws.cell(row=r, column=8).alignment = Alignment(horizontal="center")
        r += 1

    ws.auto_filter.ref = f"A4:M{r - 1}"
    last = r + 400
    for col_idx, vocab in ((5, ["programs", "apply", "entity", "feed", "listing"]),
                           (6, ["auto", "rss", "sitemap", "jsonld", "html", "skip"]),
                           (8, ["yes", "no"]),
                           (9, ["1", "2", "3"])):
        dv = DataValidation(type="list", formula1='"' + ",".join(vocab) + '"',
                            allow_blank=True, showDropDown=False)
        ws.add_data_validation(dv)
        col = get_column_letter(col_idx)
        dv.add(f"{col}5:{col}{last}")
    return ws, r - 5


# ------------------------------------------------------------------ VOCAB
def sheet_vocab(wb):
    ws = wb.create_sheet("VOCAB")
    ws.sheet_view.showGridLines = False
    banner(ws, "Controlled vocabularies",
           "The PROGRAMS and SCRAPE_CONFIG sheets validate against these lists.", 6)
    cols = [("Entity type", ENTITY_TYPES), ("Programme type", PROGRAM_TYPES),
            ("Stage", STAGES), ("Confidence", ["high", "medium", "low"]),
            ("url_role", ["programs", "apply", "entity", "feed", "listing"]),
            ("method", ["auto", "rss", "sitemap", "jsonld", "html", "skip"])]
    for i, (name, vals) in enumerate(cols, start=1):
        c = ws.cell(row=4, column=i, value=name)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = SUBHEAD
        c.border = BOX
        c.alignment = Alignment(vertical="center", indent=1)
        ws.column_dimensions[get_column_letter(i)].width = 28
        for j, v in enumerate(vals, start=5):
            cc = ws.cell(row=j, column=i, value=v)
            cc.border = BOX
            cc.font = Font(size=10)
            cc.alignment = Alignment(indent=1)
    ws.row_dimensions[4].height = 24
    return ws


# ------------------------------------------------------------------ COVERAGE
def sheet_coverage(wb, n_prog, n_ent, n_src):
    ws = wb.create_sheet("COVERAGE & GAPS")
    ws.sheet_view.showGridLines = False
    banner(ws, "Coverage & gaps", "What this file does and does not cover.", 3)
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 104

    r = 4
    c = ws.cell(row=r, column=2, value="AT A GLANCE")
    c.font = Font(bold=True, size=11, color=ACCENT)
    r += 1
    from collections import Counter
    counts = Counter(p["country"] for p in PROGRAMS)
    facts = [
        ("Programme rows", str(n_prog)),
        ("Unique entities", str(n_ent)),
        ("Markets covered", str(len(counts))),
        ("URLs queued for scraping", str(n_src)),
        ("Confidence split",
         ", ".join(f"{k}: {v}" for k, v in
                   Counter(p["confidence"] for p in PROGRAMS).most_common())),
        ("Rows by market",
         ", ".join(f"{k} {v}" for k, v in counts.most_common())),
    ]
    for label, val in facts:
        ws.cell(row=r, column=2, value=label).font = Font(bold=True, size=10)
        cc = ws.cell(row=r, column=3, value=val)
        cc.alignment = Alignment(wrap_text=True, vertical="top")
        cc.font = Font(size=10)
        ws.row_dimensions[r].height = max(15, 13 * (len(val) // 100 + 1))
        r += 1

    r += 1
    ws.cell(row=r, column=2, value="NOTES").font = Font(bold=True, size=11, color=ACCENT)
    r += 1
    for title, body in COVERAGE_NOTES:
        ws.cell(row=r, column=2, value=title).font = Font(bold=True, size=10, color=INK)
        ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        cc = ws.cell(row=r, column=3, value=body)
        cc.alignment = Alignment(wrap_text=True, vertical="top")
        cc.font = Font(size=10)
        ws.row_dimensions[r].height = max(28, 13 * (len(body) // 100 + 1))
        r += 1
    return ws


def build(out_path):
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    sheet_readme(wb)
    _, n_prog = sheet_programs(wb)
    _, n_ent = sheet_entities(wb)
    _, n_src = sheet_scrape_config(wb)
    sheet_vocab(wb)
    sheet_coverage(wb, n_prog, n_ent, n_src)
    wb.properties.created = FIXED_TIMESTAMP
    wb.properties.modified = FIXED_TIMESTAMP
    wb.save(out_path)
    normalize_zip(out_path)
    return out_path, n_prog, n_ent, n_src


if __name__ == "__main__":
    p, a, b, c = build(sys.argv[1] if len(sys.argv) > 1
                       else "output/MENA_Startup_Programs_Scout.xlsx")
    print(f"wrote {p}\n  {a} programme rows | {b} entities | {c} scrape URLs")
