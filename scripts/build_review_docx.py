"""
Render a review to Word, from the same content the PDF is built from.

The PDF and the .docx are two renderings of one `DOC` dictionary, so a change to the
words lands in both and neither can quietly fall behind the other. Formatting is
hand-built rather than styled, because a recipient's Normal style is unknown and a
document that inherits Calibri from their template is not the same document.

Inline markup is limited to what the reviews use: <b> for emphasis and the named
entity &amp;. It is parsed rather than stripped, so a bold run in a table cell stays
bold in Word.

    python3 scripts/build_review_docx.py            # every review
    python3 scripts/build_review_docx.py rv01       # one, by slug prefix
"""

import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document                                    # noqa: E402
from docx.enum.section import WD_ORIENT                      # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT               # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH                # noqa: E402
from docx.shared import Cm, Pt, RGBColor                     # noqa: E402

from scripts import _determinism as D                        # noqa: E402
from scripts import _ooxml as X                              # noqa: E402
from whitepapers import brand as B                           # noqa: E402
from whitepapers.library import REVIEWS                      # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "output", "resources")

INK = RGBColor(0x0B, 0x0E, 0x14)
MUTED = RGBColor(0x6B, 0x72, 0x85)
SECOND = RGBColor(0x3A, 0x40, 0x50)
FAINT = RGBColor(0x9A, 0xA0, 0xAF)
UV = RGBColor(0x34, 0x18, 0xE0)

FONT = "Satoshi"
FALLBACK = "Segoe UI"   # named second so Word substitutes something close, not Calibri


def runs_from(markup):
    """Split a fragment into (text, bold) pairs, honouring <b> and &amp;."""
    out = []
    for piece in re.split(r"(<b>.*?</b>)", str(markup), flags=re.S):
        if not piece:
            continue
        bold = piece.startswith("<b>")
        text = re.sub(r"</?b>", "", piece)
        text = text.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
        out.append((text, bold))
    return out


def write(par, markup, size, colour=INK, bold=False, italic=False):
    for text, is_bold in runs_from(markup):
        run = par.add_run(text)
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.color.rgb = colour
        run.bold = bold or is_bold
        run.italic = italic
        X.run_cs(run, FALLBACK, size)
    return par


def para(doc, markup, size=9.4, colour=INK, space_after=2.4, bold=False,
         space_before=0, line=1.5):
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.space_after = Pt(space_after * 2.835)
    fmt.space_before = Pt(space_before * 2.835)
    fmt.line_spacing = line
    write(p, markup, size, colour, bold)
    return p


def build_one(doc_spec):
    document = Document()

    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)

    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(9.4)
    normal.font.color.rgb = INK

    # ------------------------------------------------------------- masthead
    head = document.add_paragraph()
    head.paragraph_format.space_after = Pt(1)
    X.para_border_top(head, "3418E0", 24)
    logo = head.add_run()
    logo.add_picture(os.path.join(ROOT, B.LOGO_INK), width=Cm(3.0))

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    meta.paragraph_format.space_after = Pt(6)
    write(meta, f"{doc_spec['series'].upper()} · {doc_spec['number'].upper()}\n"
                f"{doc_spec['cover_foot'].upper()}", 7.4, FAINT, bold=True)

    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    write(title, doc_spec["title"], 20, INK, bold=True)

    para(document, doc_spec["standfirst"], 9.6, SECOND, space_after=4)

    # ---------------------------------------------------------------- blocks
    for blk in doc_spec["blocks"]:
        kind = blk[0]

        if kind == "p":
            para(document, blk[1])

        elif kind == "h2":
            para(document, blk[1], 12, INK, bold=True, space_before=3, space_after=1.5)

        elif kind == "statement":
            para(document, blk[1], 16, INK, bold=True, space_after=3, line=1.15)

        elif kind == "table":
            headers, rows = blk[1], blk[2]
            table = document.add_table(rows=1, cols=len(headers))
            table.alignment = WD_TABLE_ALIGNMENT.LEFT
            table.autofit = False
            widths = [Cm(1.0), Cm(4.6), Cm(2.0), Cm(9.0)][:len(headers)]
            for cell, text, width in zip(table.rows[0].cells, headers, widths):
                cell.width = width
                cell.paragraphs[0].paragraph_format.space_after = Pt(3)
                write(cell.paragraphs[0], str(text).upper(), 7.4, SECOND, bold=True)
                X.tc_borders_bottom(cell, "0B0E14", 12)
            for row in rows:
                cells = table.add_row().cells
                for i, (cell, value, width) in enumerate(zip(cells, row, widths)):
                    cell.width = width
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.line_spacing = 1.3
                    write(p, value, 8.4, INK if i == 0 else SECOND)
                    X.tc_borders_bottom(cell, "E2E4EA", 6)

        elif kind == "num":
            for number, heading, body in blk[1]:
                p = document.add_paragraph()
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.space_before = Pt(5)
                write(p, f"{number}   ", 11, UV, bold=True)
                write(p, heading, 9.4, INK, bold=True)
                b = document.add_paragraph()
                b.paragraph_format.left_indent = Cm(1.0)
                b.paragraph_format.space_after = Pt(2)
                b.paragraph_format.line_spacing = 1.4
                write(b, body, 8.8, SECOND)

        elif kind == "callout":
            p = document.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(1)
            write(p, blk[1], 9.4, INK, bold=True)
            para(document, blk[2], 8.8, SECOND, space_after=3)

        elif kind == "note":
            p = document.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.4
            X.para_border_top(p, "E2E4EA", 6)
            write(p, blk[1], 7.9, MUTED)

        elif kind in ("break", "rule"):
            continue

        else:
            raise ValueError(f"block not supported in the Word renderer: {kind}")

    # ------------------------------------------------------- about StartPad
    band = document.add_paragraph()
    band.paragraph_format.space_before = Pt(10)
    band.paragraph_format.space_after = Pt(2)
    X.para_border_top(band, "0B0E14", 18)
    mark = band.add_run()
    mark.add_picture(os.path.join(ROOT, B.LOGO_INK), width=Cm(2.8))

    about = document.add_paragraph()
    about.paragraph_format.space_after = Pt(1)
    about.paragraph_format.line_spacing = 1.4
    write(about, "StartPad", 7.9, INK, bold=True)
    write(about, " is a founder-readiness platform for Egypt and MENA. It moves a "
                 "young person with an idea and no company through fifteen structured "
                 "missions and issues verifiable proof of the work. Nothing submitted "
                 "is taken from the person who submitted it — it is sealed and "
                 "timestamped to them.", 7.9, MUTED)

    url = document.add_paragraph()
    url.paragraph_format.space_after = Pt(0)
    write(url, B.URL, 11, INK, bold=True)

    document.core_properties.title = doc_spec["title"]
    document.core_properties.author = "StartPad"
    document.core_properties.comments = doc_spec["standfirst_plain"]
    document.core_properties.keywords = ", ".join(doc_spec["keywords"])
    document.core_properties.last_modified_by = "StartPad"

    path = os.path.join(OUT, f"{doc_spec['file']}.docx")
    document.save(path)
    D.normalize_zip(path)
    return path


def main(argv):
    os.makedirs(OUT, exist_ok=True)
    wanted = [a.lower() for a in argv[1:]]
    docs = [d for d in REVIEWS
            if not wanted or any(d["slug"].startswith(w) for w in wanted)]
    if not docs:
        print("no reviews matched", wanted)
        return 1
    for spec in docs:
        path = build_one(spec)
        print(f"  {os.path.basename(path):<52} {os.path.getsize(path) / 1024:>6.0f} KB")
    print(f"\n{len(docs)} Word file(s) written to output/resources/")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
