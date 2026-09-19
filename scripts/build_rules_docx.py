# -*- coding: utf-8 -*-
"""
Build the bilingual Mission Zero rules .docx from competition/rules.py.

Arabic first, then English, each with its own direction. Tables need one extra
signal beyond the paragraph-level RTL used elsewhere: `w:bidiVisual` on the table
itself, which reverses column order. Without it an Arabic table reads its columns
left-to-right while its text reads right-to-left, which is worse than either.
"""
import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from competition.rules import VERSIONS  # noqa: E402
from scripts import _ooxml as X  # noqa: E402
from scripts._determinism import FIXED_TIMESTAMP, normalize_zip  # noqa: E402

INK = RGBColor(0x10, 0x1E, 0x2E)
SEA = RGBColor(0x0E, 0x7C, 0x86)
AMBER = RGBColor(0xB4, 0x65, 0x0E)
GREY = RGBColor(0x63, 0x73, 0x7F)
ARABIC_FONT = "Arial"


def _bidi_para(p):
    X.para_rtl(p)


def _bidi_table(table):
    X.tbl_rtl(table)


def _style_run(run, size, rtl, bold=False, italic=False, color=INK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    X.run_cs(run, ARABIC_FONT if rtl else "Calibri", size, rtl)


def para(doc, text="", size=10.5, rtl=False, bold=False, italic=False, color=INK,
         before=0, after=6, indent=0, align=None, line=1.25):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if indent:
        (pf.__setattr__("right_indent", Cm(indent)) if rtl
         else pf.__setattr__("left_indent", Cm(indent)))
    if rtl:
        _bidi_para(p)
        p.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.RIGHT
    else:
        p.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.LEFT
    if text:
        _style_run(p.add_run(text), size, rtl, bold, italic, color)
    return p


def shade(cell, hexcolor):
    X.tc_shade(cell, hexcolor)


def borders(table, color="D9DBD3"):
    X.tbl_borders(table, color=color, sz="4")


def cell_text(cell, text, rtl, size=9, bold=False, color=INK):
    cell.text = ""
    first = True
    for line in str(text).split("\n"):
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.15
        if rtl:
            _bidi_para(p)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _style_run(p.add_run(line), size, rtl, bold, color=color)


def table(doc, headers, rows, rtl):
    t = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    if rtl:
        _bidi_table(t)
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        shade(c, "EDF3F3")
        cell_text(c, h, rtl, size=8.5, bold=True, color=SEA)
    for r, row in enumerate(rows, start=1):
        for i, val in enumerate(row):
            cell_text(t.rows[r].cells[i], val, rtl)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def render(doc, v):
    rtl = v["dir"] == "rtl"
    para(doc, v["name"], size=24, rtl=rtl, bold=True, color=INK, after=2, line=1.1)
    para(doc, v["tagline"], size=12, rtl=rtl, italic=True, color=GREY, after=8)
    para(doc, v["doc_title"], size=11, rtl=rtl, bold=True, color=SEA, after=2)
    para(doc, v["version_line"], size=8.5, rtl=rtl, color=GREY, after=14)

    for block in v["blocks"]:
        kind = block[0]
        if kind == "h":
            para(doc, block[1], size=13.5, rtl=rtl, bold=True, color=SEA,
                 before=14, after=5)
        elif kind == "p":
            para(doc, block[1], size=10.5, rtl=rtl, after=7)
        elif kind == "b":
            for item in block[1]:
                para(doc, "•  " + item, size=10, rtl=rtl, indent=0.6, after=3)
        elif kind == "n":
            for i, item in enumerate(block[1], 1):
                para(doc, f"{i}.  " + item, size=10, rtl=rtl, indent=0.6, after=3)
        elif kind == "t":
            table(doc, block[1], block[2], rtl)
        elif kind == "cal":
            table(doc,
                  (["متى", "ماذا", "التفاصيل"] if rtl else ["When", "What", "Detail"]),
                  [list(r) for r in block[1]], rtl)


def build(out_path):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.font.color.rgb = INK
    for s in doc.sections:
        s.top_margin = Cm(1.9)
        s.bottom_margin = Cm(1.9)
        s.left_margin = Cm(2.0)
        s.right_margin = Cm(2.0)

    for i, v in enumerate(VERSIONS):
        if i:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        render(doc, v)

    core = doc.core_properties
    core.title = "Mission Zero — rules and judging criteria"
    core.author = "StartPad"
    core.created = FIXED_TIMESTAMP
    core.modified = FIXED_TIMESTAMP

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    doc.save(out_path)
    normalize_zip(out_path)
    return out_path


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "output/Mission_Zero_Rules_AR_EN.docx"
    p = build(out)
    print("wrote", p, os.path.getsize(p), "bytes")
