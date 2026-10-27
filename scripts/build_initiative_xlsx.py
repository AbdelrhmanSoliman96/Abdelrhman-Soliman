"""Build the sponsorship invitation list for One Million First-Time Founders."""
import math
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from outreach import initiative as I                            # noqa: E402
from scripts._determinism import FIXED_TIMESTAMP, normalize_zip  # noqa: E402
from scripts.build_prospects_xlsx import (                      # noqa: E402
    BOX, CONF_FILL, CONF_FONT, HEAD, INK, ACCENT, SUNK, SUBHEAD, banner,
    header, prose,
)

ROLE_FILL = {
    "CSR sponsor": PatternFill("solid", fgColor="0F5C4B"),
    "NGO and development partner": PatternFill("solid", fgColor="2E6F5E"),
    "Community supporter": PatternFill("solid", fgColor="7FA99B"),
}
ROLE_FONT = Font(size=9.5, bold=True, color="FFFFFF")

STAGES = ["Not invited", "Invited", "Replied", "Meeting held",
          "Proposal sent", "Negotiating", "Signed", "Declined", "Parked"]


def money(v, dp=0):
    return f"${v:,.{dp}f}"


# ------------------------------------------------------------------ READ ME
def sheet_readme(wb):
    ws = wb.create_sheet("READ ME FIRST")
    ws.sheet_view.showGridLines = False
    rows = I.invitation_list()
    f = dict((k, v) for k, v, _ in I.funnel())
    banner(ws, f"{I.NAME}: who to invite, and as what",
           f"{len(rows)} organisations in three roles. {I.WAVE}.", 5)
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 108

    counts = {}
    for r in rows:
        counts[r["role"]] = counts.get(r["role"], 0) + 1

    blocks = [
        ("WHAT THIS IS", ""),
        ("", f"The invitation list for one named initiative, cut from the "
             f"wider prospect registry in "
             f"StartPad_Egypt_CSR_NGO_Prospects.xlsx. That file asks who in "
             f"Egypt could fund entrepreneurship work at all. This one asks "
             f"who we invite to sponsor {I.NAME}, in which of three roles, "
             f"for what contribution, and what they get back."),
        ("", "The framework document is the companion to this file: "
             "StartPad_Initiative_Framework.docx. Send the framework, work "
             "the list."),
        ("", ""),
        ("THE THREE ROLES", ""),
    ]
    for name, r in I.ROLES.items():
        blocks.append((f"{name} ({counts.get(name, 0)})",
                       f"{r['who']}  GIVES: {r['gives']}  GETS: {r['gets']}"))
    blocks += [
        ("", ""),
        ("WHO IS NOT ON THIS LIST, AND WHY", ""),
    ]
    for cat, why in I.NOT_INVITED.items():
        blocks.append((cat, why))
    blocks += [
        ("", ""),
        ("THE NUMBER THIS LIST HAS TO PRODUCE", ""),
        ("", f"The wave costs {money(I.budget_total())}, which is EGP "
             f"{I.egp(I.budget_total()):,.0f}, and it is funded entirely by "
             f"the organisations on this list. StartPad puts in no cash at "
             f"all and takes no margin."),
        ("", f"The three cash tiers raise {money(I.target_money())} if "
             f"every slot sells, so the coalition can lose four or five "
             f"conversations and still fund the wave."),
        ("", f"At that budget the initiative registers "
             f"{f['Registered entrepreneurs']:,.0f} entrepreneurs at "
             f"{money(I.budget_total() / f['Registered entrepreneurs'], 2)} "
             f"a head, puts {f['All 15 missions completed']:,.0f} of them "
             f"through all fifteen missions, and leaves "
             f"{f['Ideas in the bank']:,.0f} ideas in the bank. Everything "
             f"is online and free, there are no paid cohorts, and offline "
             f"sessions happen only where a partner asks for one and hosts "
             f"it."),
        ("", "The full budget and what StartPad gets back are in the "
             "companion workbook: StartPad_Initiative_Budget.xlsx."),
        ("", ""),
        ("HOW TO WORK IT", ""),
        ("", "Founding Partners first, and do not announce until both "
             "have signed. A coalition with nobody in it is a hard sell; a "
             "coalition with two recognisable names in it sells itself to "
             "the next fifteen."),
        ("", "Community supporters second, and before the summer, because "
             "student leadership turns over in July and August. They cost "
             "nothing, they are the highest converting reach in the whole "
             "plan, and the referral loop runs on them."),
        ("", "NGO and development partners third. They bring reach, "
             "governorates and the venues for offline info sessions rather "
             "than money, and several will want to co-apply for donor "
             "funding, which is a slower and larger conversation."),
        ("", ""),
        ("THE COLUMNS YOU FILL IN", ""),
        ("", "Stage, owner, pledged amount, next action, next date and "
             "notes. Pledged amount is what makes the COVERAGE sheet mean "
             "anything, so fill it in even when it is a guess."),
        ("", ""),
        ("CONFIDENCE, UNCHANGED FROM THE REGISTRY", ""),
        ("high", "Organisation and focus confirmed in published sources."),
        ("medium", "Established organisation; the specific programme or "
                   "route changes between cycles and must be re-checked."),
        ("low", "A reasoned prospect. Verify before spending a meeting."),
        ("", "No individual's name, email address or phone number appears "
             "anywhere in this file, for the same reason as in the "
             "registry: a stale contact burns the introduction."),
    ]
    prose(ws, blocks)
    return ws


# --------------------------------------------------------- INVITATION LIST
COLS = [
    ("role", "Role", 24), ("tier", "Invite at", 20),
    ("name", "Organisation", 34), ("arabic", "Arabic name", 20),
    ("sector", "Sector", 22), ("reach", "Reach", 22),
    ("contribution", "What we ask them for", 28),
    ("offer", "What they get", 50),
    ("focus", "What they already fund", 52),
    ("want", "What they are looking for", 44),
    ("unit", "Who decides", 34),
    ("route", "How to reach them", 48),
    ("window", "When to approach", 30),
    ("tier_abc", "Priority", 9),
    ("conf", "Confidence", 11),
    ("note", "Verify before quoting", 44),
]
OWN = [("Stage", 16), ("Owner", 14), ("Pledged USD", 13),
       ("Next action", 30), ("Next date", 13), ("Notes", 34)]
WRAP = {"name", "arabic", "sector", "reach", "contribution", "offer", "focus",
        "want", "unit", "route", "window", "note", "role", "tier"}


def sheet_list(wb):
    ws = wb.create_sheet("INVITATION LIST")
    ws.sheet_view.showGridLines = False
    heads = [c[1] for c in COLS] + [o[0] for o in OWN]
    widths = [c[2] for c in COLS] + [o[1] for o in OWN]
    header(ws, 1, heads, widths, freeze_col=4)
    width_of = dict((c[0], c[2]) for c in COLS)

    r = 2
    for p in I.invitation_list():
        for i, (key, _, _) in enumerate(COLS, start=1):
            c = ws.cell(row=r, column=i, value=p[key])
            c.border = BOX
            c.font = Font(size=10)
            c.alignment = Alignment(wrap_text=key in WRAP, vertical="top",
                                    indent=1)
        rc = ws.cell(row=r, column=1)
        rc.fill = ROLE_FILL[p["role"]]
        rc.font = ROLE_FONT
        ws.cell(row=r, column=2).font = Font(size=10, bold=True, color=ACCENT)
        ws.cell(row=r, column=3).font = Font(size=10, bold=True, color=INK)
        if p["arabic"]:
            ws.cell(row=r, column=4).alignment = Alignment(
                horizontal="right", vertical="top", wrap_text=True)
        pr = ws.cell(row=r, column=14)
        pr.alignment = Alignment(horizontal="center", vertical="top")
        pr.font = Font(size=10, bold=True, color=INK)
        cf = ws.cell(row=r, column=15)
        cf.fill = CONF_FILL[p["conf"]]
        cf.font = CONF_FONT[p["conf"]]
        cf.alignment = Alignment(horizontal="center", vertical="top")
        for i in range(len(COLS) + 1, len(COLS) + len(OWN) + 1):
            cell = ws.cell(row=r, column=i)
            cell.border = BOX
            cell.fill = SUNK
            cell.alignment = Alignment(wrap_text=True, vertical="top",
                                       indent=1)
        s = ws.cell(row=r, column=len(COLS) + 1, value="Not invited")
        s.font = Font(size=10, color="6B7A74")
        ws.cell(row=r, column=len(COLS) + 3).number_format = "#,##0"
        need = max((len(str(p[k])) / width_of[k] for k in WRAP if p[k]),
                   default=1)
        ws.row_dimensions[r].height = max(30, 12.2 * (math.ceil(need * 1.12)
                                                      + 1))
        r += 1

    last = get_column_letter(len(COLS) + len(OWN))
    ws.auto_filter.ref = f"A1:{last}{r - 1}"
    col = get_column_letter(len(COLS) + 1)
    dv = DataValidation(type="list", formula1='"' + ",".join(STAGES) + '"',
                        allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{r - 1}")
    return ws, r - 1


# ---------------------------------------------------------------- PACKAGES
def sheet_tiers(wb):
    ws = wb.create_sheet("PACKAGES")
    ws.sheet_view.showGridLines = False
    banner(ws, "What a partner gets for what they give",
           "Slots are capped on purpose. A tier anybody can buy is not "
           "worth buying.", 5)
    for col, w in zip("ABCDE", [3, 26, 14, 10, 96]):
        ws.column_dimensions[col].width = w
    r = 4
    for t in I.TIERS:
        c = ws.cell(row=r, column=2, value=t["name"])
        c.font = Font(size=12, bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(vertical="center", indent=1)
        price = money(t["price"]) if t["price"] else "No cash"
        p = ws.cell(row=r, column=3, value=price)
        p.font = Font(size=12, bold=True, color="FFFFFF")
        p.fill = HEAD
        p.alignment = Alignment(horizontal="center", vertical="center")
        slot = f"{t['slots']} slots" if t["slots"] else "Uncapped"
        sc = ws.cell(row=r, column=4, value=slot)
        sc.font = Font(size=9.5, bold=True, color="FFFFFF")
        sc.fill = HEAD
        sc.alignment = Alignment(horizontal="center", vertical="center")
        g = ws.cell(row=r, column=5, value=t["gives"])
        g.font = Font(size=9.5, italic=True, color="FFFFFF")
        g.fill = HEAD
        g.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
        ws.row_dimensions[r].height = max(24, 12.4 * (len(t["gives"]) // 92
                                                      + 1))
        r += 1
        for b in t["benefits"]:
            ws.cell(row=r, column=2, value=t["role"]).font = Font(
                size=8.5, color="6B7A74")
            ws.cell(row=r, column=2).alignment = Alignment(
                vertical="top", wrap_text=True, indent=1)
            c = ws.cell(row=r, column=5, value=b)
            c.font = Font(size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
            c.border = BOX
            ws.row_dimensions[r].height = max(16, 12.4 * (len(b) // 92 + 1))
            r += 1
        r += 1
    return ws


# ----------------------------------------------------------- IN-KIND MENU
def sheet_inkind(wb):
    ws = wb.create_sheet("IN-KIND MENU")
    ws.sheet_view.showGridLines = False
    header(ws, 1, ["What we accept", "Who has it", "Why it matters",
                   "Value to the wave", "Offered by"],
           [30, 34, 78, 16, 26])
    r = 2
    for what, who, why, val in I.INKIND:
        for i, v in enumerate([what, who, why, val], start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = BOX
            c.font = Font(size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
        ws.cell(row=r, column=1).font = Font(size=10, bold=True, color=ACCENT)
        v = ws.cell(row=r, column=4)
        v.alignment = Alignment(horizontal="center", vertical="top")
        v.fill = {"High": CONF_FILL["high"], "Medium": CONF_FILL["medium"],
                  "Low": CONF_FILL["low"]}[val]
        v.font = {"High": CONF_FONT["high"], "Medium": CONF_FONT["medium"],
                  "Low": CONF_FONT["low"]}[val]
        o = ws.cell(row=r, column=5)
        o.border = BOX
        o.fill = SUNK
        ws.row_dimensions[r].height = max(28, 12.4 * (len(why) // 76 + 1))
        r += 1
    ws.cell(row=r + 1, column=1,
            value="In-kind is valued at its market or rate card value and "
                  "credited at the cash tier it matches. A broadcaster "
                  "giving airtime worth $75,000 is a Lead Partner with every "
                  "Lead Partner benefit. This is usually faster than a cash "
                  "approval, and for a media owner or a venue owner it is "
                  "the only realistic route.").font = Font(
        size=10, italic=True, color=INK)
    ws.cell(row=r + 1, column=1).alignment = Alignment(wrap_text=True,
                                                       vertical="top")
    ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1,
                   end_column=5)
    ws.row_dimensions[r + 1].height = 52
    return ws


# ------------------------------------------------------------ WHY THEY JOIN
def sheet_values(wb):
    ws = wb.create_sheet("WHY THEY JOIN")
    ws.sheet_view.showGridLines = False
    banner(ws, "What each kind of partner is actually buying",
           "Nobody on the invitation list wants to fund entrepreneurship "
           "training. They each want their own outcome.", 4)
    ws.column_dimensions["B"].width = 44
    ws.column_dimensions["C"].width = 96
    blocks = []
    for role in I.ROLES:
        blocks.append((role.upper(), ""))
        for head, body in I.VALUES[role]:
            blocks.append((head, body))
        blocks.append(("", ""))
    blocks.append(("WHAT EVERY PARTNER GETS, WHATEVER THE ROLE", ""))
    for head, body in I.SHARED_VALUE:
        blocks.append((head, body))
    blocks.append(("", ""))
    blocks.append(("THE IDEAS BANK, AND ITS RULES", ""))
    for head, body in I.IDEAS_BANK:
        blocks.append((head, body))
    prose(ws, blocks, wrap_at=93)
    return ws


# ------------------------------------------------------------------ COVERAGE
def sheet_coverage(wb, list_rows):
    ws = wb.create_sheet("COVERAGE")
    ws.sheet_view.showGridLines = False
    banner(ws, "Enough candidates for the slots?",
           "Pledged amounts come from the INVITATION LIST sheet, so this "
           "sheet is live once you start filling it in.", 7)
    for col, w in zip("ABCDEFG", [3, 26, 12, 14, 14, 16, 18]):
        ws.column_dimensions[col].width = w
    r = 4
    for i, h in enumerate(["Tier", "Slots", "Candidates", "Cover",
                           "Slot value", "Tier target"], start=2):
        c = ws.cell(row=r, column=i, value=h)
        c.font = Font(bold=True, size=10, color="FFFFFF")
        c.fill = SUBHEAD
        c.border = BOX
        c.alignment = Alignment(vertical="center", indent=1, wrap_text=True)
    ws.row_dimensions[r].height = 28
    r += 1
    first = r
    for name, slots, cands, price, target in I.coverage():
        ws.cell(row=r, column=2, value=name).font = Font(size=10, bold=True)
        ws.cell(row=r, column=3, value=slots)
        ws.cell(row=r, column=4, value=cands)
        cov = ws.cell(row=r, column=5, value=f"=D{r}/C{r}")
        cov.number_format = "0.0x"
        ws.cell(row=r, column=6, value=price).number_format = "#,##0"
        ws.cell(row=r, column=7, value=target).number_format = "#,##0"
        for i in range(2, 8):
            ws.cell(row=r, column=i).border = BOX
            if i > 2:
                ws.cell(row=r, column=i).alignment = Alignment(
                    horizontal="center")
        if cands < slots:
            for i in (4, 5):
                ws.cell(row=r, column=i).fill = CONF_FILL["low"]
                ws.cell(row=r, column=i).font = CONF_FONT["low"]
        r += 1
    for i, col in zip(range(3, 8), "CDEFG"):
        if col in ("E", "F"):
            continue
        c = ws.cell(row=r, column=i,
                    value=f"=SUM({col}{first}:{col}{r-1})")
        c.font = Font(size=10, bold=True, color="FFFFFF")
        c.fill = HEAD
        c.border = BOX
        c.alignment = Alignment(horizontal="center")
        c.number_format = "#,##0"
    t = ws.cell(row=r, column=2, value="TOTAL")
    t.font = Font(size=10, bold=True, color="FFFFFF")
    t.fill = HEAD
    t.alignment = Alignment(indent=1)
    r += 2

    pledged_col = get_column_letter(len(COLS) + 3)
    stage_col = get_column_letter(len(COLS) + 1)
    rows = [
        ("Organisations invited",
         f'=COUNTIF(\'INVITATION LIST\'!{stage_col}2:{stage_col}{list_rows},'
         f'"<>Not invited")', "#,##0"),
        ("Signed",
         f'=COUNTIF(\'INVITATION LIST\'!{stage_col}2:{stage_col}{list_rows},'
         f'"Signed")', "#,##0"),
        ("Pledged, USD",
         f"=SUM('INVITATION LIST'!{pledged_col}2:{pledged_col}{list_rows})",
         "#,##0"),
        ("Wave 1 budget needed",
         f"={I.target_money() * I.TARGET_FILL}", "#,##0"),
        ("Gap",
         f"={I.target_money() * I.TARGET_FILL}"
         f"-SUM('INVITATION LIST'!{pledged_col}2:{pledged_col}{list_rows})",
         "#,##0"),
    ]
    ws.cell(row=r, column=2, value="WHERE THE RAISE STANDS").font = Font(
        bold=True, size=11, color=ACCENT)
    r += 1
    for label, formula, fmt in rows:
        ws.cell(row=r, column=2, value=label).font = Font(size=10, bold=True)
        c = ws.cell(row=r, column=3, value=formula)
        c.number_format = fmt
        c.font = Font(size=11, bold=True, color=ACCENT)
        c.fill = SUNK
        c.border = BOX
        c.alignment = Alignment(horizontal="center")
        r += 1
    r += 1
    thin = [n for n, slots, cands, _, _ in I.coverage() if cands < slots]
    note = ("Every tier has more candidates than slots. Work the list in "
            "priority order and stop when a tier fills.")
    if thin:
        note = ("Under-covered tiers: " + ", ".join(thin) + ". There are "
                "fewer candidates sized for these than there are slots, so "
                "either the slot count comes down or mid-sized Egyptian "
                "companies have to be added to the registry. The Mission "
                "Partner tier at $12,000 is the one a mid-cap can sign "
                "inside a quarter, and the registry is thin on mid-caps "
                "because it was built from the organisations that publish "
                "their CSR.")
    ws.cell(row=r, column=2, value=note).font = Font(size=10, italic=True,
                                                     color=INK)
    ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True,
                                                   vertical="top")
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
    ws.row_dimensions[r].height = 60
    return ws


def build(path):
    wb = Workbook()
    wb.remove(wb.active)
    sheet_readme(wb)
    _, list_rows = sheet_list(wb)
    sheet_tiers(wb)
    sheet_inkind(wb)
    sheet_values(wb)
    sheet_coverage(wb, list_rows)
    wb.properties.creator = "StartPad"
    wb.properties.title = f"{I.NAME}: sponsorship invitation list"
    wb.properties.created = FIXED_TIMESTAMP
    wb.properties.modified = FIXED_TIMESTAMP
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    normalize_zip(path)
    return path


if __name__ == "__main__":
    out = (sys.argv[1] if len(sys.argv) > 1 else
           os.path.join(os.path.dirname(os.path.dirname(
               os.path.abspath(__file__))), "output",
               "StartPad_Initiative_Sponsor_List.xlsx"))
    build(out)
    rows = I.invitation_list()
    print(f"wrote {out}  ({os.path.getsize(out)/1024:,.0f} KB)")
    print(f"{len(rows)} organisations invited, in {len(I.ROLES)} roles")
