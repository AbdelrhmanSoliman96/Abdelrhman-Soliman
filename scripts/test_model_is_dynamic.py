"""
Prove the financial model is dynamic, by moving inputs and watching.

A model can contain nothing but formulas and still be dead, if those formulas
point at the wrong cells. The only convincing test is to change an assumption
and check that every sheet downstream of it moves, and that the balance sheet
still balances afterwards.

Nine inputs are moved at once, chosen to enter the model at nine different
points: a price, a volume, a cost per click, a conversion rate, a facilitator
day rate, a salary, a headcount, a collection period and a revenue share. Then
the whole workbook is recalculated and the outputs are compared.

    python scripts/test_model_is_dynamic.py output/StartPad_Financial_Model.xlsx
"""
import os
import shutil
import subprocess
import sys

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))


def find(ws, label, col=1, nth=1):
    """Row of the nth cell in `col` whose text is exactly `label`."""
    seen = 0
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, col).value
        if isinstance(v, str) and v.strip() == label:
            seen += 1
            if seen == nth:
                return r
    raise KeyError(f"{label} (occurrence {nth})")


# (sheet, row label, column, multiplier, what it should move)
PERTURB = [
    ("Assumption Setting", "University cohort, all in programme fee", 3, 1.40,
     "a price"),
    ("Assumption Setting", "University cohorts, Egypt, paid", 4, 3.00,
     "a volume"),
    ("Assumption Setting", "Short video, TikTok, Instagram, Facebook", 8, 2.00,
     "a cost per click"),
    ("Assumption Setting", "Stage 3, registered to paid, reached online", 3,
     2.50, "a conversion rate"),
    ("Assumption Setting", "Facilitator day rate", 3, 1.60,
     "a direct cost driver"),
    ("Assumption Setting", "Days sales outstanding, institutional", 3, 1.80,
     "a collection period"),
    ("Assumption Setting", "Party 1, the partner's share of that cohort fee",
     3, 1.50, "a revenue share"),
    ("Payroll", "Chief Executive Officer", 4, 2.00, "a salary"),
    ("Payroll", "Senior full-stack engineer", 9, 3.00, "a headcount"),
]

# (sheet, row label, column) of outputs that must respond
# (sheet, row label, column, [occurrence of the label])
WATCH = [
    ("Total Revenue", "Total revenue", 2),
    ("Revenue Streams", "Institutional and sponsor revenue", 2),
    ("Revenue Streams", "Subscription revenue recognised", 15),
    ("Acquisitions", "Total clicks", 2),
    ("Acquisitions", "Stage 3, new paid accounts", 2),
    ("Acquisitions", "Total contracts signed", 15),
    ("Expenses", "Total direct cost", 2),
    ("Expenses", "Revenue share, certified delivery partners", 15),
    ("Expenses", "Total operating cost", 2),
    ("Payroll", "Total payroll", 2),
    ("Capital Expenditure", "Total capital expenditure", 2),
    ("Depreciation and Amortization", "Total depreciation and amortisation",
     15),
    ("Income Statement", "Total revenue", 3),
    ("Income Statement", "Net profit for the year", 5),
    ("Cash Flow Statement", "Trade receivables", 15),
    ("Cash Flow Statement", "Net cash from operating activities", 2),
    ("Balance Sheet", "Total assets", 5),
    ("Financial Indicators", "EBITDA", 3),
    ("Financial Indicators", "Cost per paid account", 3),
    ("Financial Indicators", "Coverage, signings against the plan", 4),
    ("BEP Analysis", "Break even revenue", 3),
    ("PB Analysis", "Cumulative free cash flow", 5),
    ("Stress Test", "Downside", 4, 3),
    ("Sensitivity", "Volume change", 5),
    ("Exit Angel and VC", "Proceeds at exit", 3),
    ("Summary", "EBITDA", 4),
]

CHECKS = [("Balance Sheet", "CHECK, total assets less liabilities and equity")]


def read(path):
    wb = load_workbook(path, data_only=True)
    wf = load_workbook(path)
    out = {}
    for entry in WATCH:
        sheet, label, col = entry[:3]
        nth = entry[3] if len(entry) > 3 else 1
        r = find(wf[sheet], label, nth=nth)
        out[(sheet, label, col)] = wb[sheet].cell(r, col).value
    checks = {}
    for sheet, label in CHECKS:
        r = find(wf[sheet], label)
        checks[(sheet, label)] = [wb[sheet].cell(r, c).value
                                  for c in range(3, 9)]
    return out, checks


def recalc(path):
    subprocess.run([sys.executable, os.path.join(HERE, "recalc.py"), path,
                    "900"], check=True, capture_output=True)


def main(path):
    work = path.replace(".xlsx", "_perturbed.xlsx")
    shutil.copy(path, work)
    base, base_checks = read(path)

    wf = load_workbook(work)
    print("MOVING NINE INPUTS")
    for sheet, label, col, mult, what in PERTURB:
        ws = wf[sheet]
        r = find(ws, label)
        old = ws.cell(r, col).value
        ws.cell(r, col).value = old * mult
        print(f"  {what:24s} {sheet}!{ws.cell(r, col).coordinate}  "
              f"{old} -> {old * mult}")
    wf.save(work)
    recalc(work)
    after, after_checks = read(work)

    print("\nWHAT MOVED")
    stuck = []
    for key in base:
        b, a = base[key], after[key]
        b = 0 if b is None else b
        a = 0 if a is None else a
        moved = abs(a - b) > max(1e-9, abs(b) * 1e-9)
        sheet, label, col = key
        mark = "moved" if moved else "STUCK"
        if not moved:
            stuck.append(f"{sheet}!{label}")
        print(f"  {mark}  {sheet:30s} {label[:40]:42s} "
              f"{b:>14,.2f} -> {a:>14,.2f}")

    print("\nBALANCE SHEET STILL BALANCES")
    broken = []
    for key, vals in after_checks.items():
        worst = max(abs(v or 0) for v in vals)
        print(f"  {key[1]}: largest absolute value {worst:,.6f}")
        if worst > 0.01:
            broken.append(key)

    os.remove(work)
    if stuck:
        print(f"\nFAILED: {len(stuck)} outputs did not respond")
        for s in stuck:
            print("  ", s)
        return 1
    if broken:
        print("\nFAILED: the balance sheet stopped balancing")
        return 1
    print(f"\nPASSED: all {len(base)} watched outputs responded and the "
          f"balance sheet still checks to nil.")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1
                  else "output/StartPad_Financial_Model.xlsx"))
