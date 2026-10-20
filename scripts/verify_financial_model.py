"""
Verify the financial model independently.

The audit proves there are no static values. The dynamic test proves the cells
are connected. Neither proves the cells are connected to the RIGHT things: a
formula can point one row off and still be a formula, still recalculate, and
still be wrong.

So this recomputes the model a second time, in Python, from the plan and from
the input values read back out of the workbook, and compares. Where the two
disagree, one of them is wrong, and the difference says where to look.

    python scripts/verify_financial_model.py output/StartPad_Financial_Model.xlsx
"""
import os
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.dirname(HERE))

import _fm as F                                                 # noqa: E402
from finmodel import plan as P                                  # noqa: E402

TOL = 0.5          # dollars
RTOL = 1e-7

results = []


def check(name, expected, actual, tol=TOL):
    ok = abs((expected or 0) - (actual or 0)) <= max(
        tol, abs(expected or 0) * RTOL)
    results.append((ok, name, expected, actual))
    return ok


class Sheet:
    """Read a sheet by row label, with both values and formulas available."""

    def __init__(self, wv, wf, name):
        self.v, self.f = wv[name], wf[name]
        self.rows = {}
        for r in range(1, self.f.max_row + 1):
            t = self.f.cell(r, 1).value
            if isinstance(t, str) and t.strip():
                self.rows.setdefault(t.strip(), []).append(r)

    def row(self, label, nth=0):
        return self.rows[label][nth]

    def m(self, label, month, nth=0):
        return self.v.cell(self.row(label, nth), F.mcol(month)).value or 0

    def y(self, label, year, nth=0):
        return self.v.cell(self.row(label, nth), F.ycol(year)).value or 0

    def c(self, label, col, nth=0):
        return self.v.cell(self.row(label, nth), col).value or 0


def main(path):
    wv = load_workbook(path, data_only=True)
    wf = load_workbook(path)
    A = Sheet(wv, wf, "Assumption Setting")
    REV = Sheet(wv, wf, "Revenue Streams")
    TOT = Sheet(wv, wf, "Total Revenue")
    ACQ = Sheet(wv, wf, "Acquisitions")
    EXP = Sheet(wv, wf, "Expenses")
    PAY = Sheet(wv, wf, "Payroll")
    DEP = Sheet(wv, wf, "Depreciation and Amortization")
    CAP = Sheet(wv, wf, "Capital Expenditure")
    IS = Sheet(wv, wf, "Income Statement")
    BS = Sheet(wv, wf, "Balance Sheet")
    CF = Sheet(wv, wf, "Cash Flow Statement")

    def ass(label, col=3):
        return A.c(label, col)

    FUNCTIONS = list(dict.fromkeys(fn for fn, *_ in P.ROLES))
    BUYER_LABELS = {
        "inst": "Institutions, universities and schools",
        "spon": "Sponsors, banks, telcos and corporates",
        "ngo": "NGOs, donors and development agencies",
        "govt": "Government and national initiatives",
        "csr": "CSR and corporate innovation budgets",
    }

    # ---------------------------------------------------------- volumes
    plan_labels = {
        "uni_cohorts_eg": "University cohorts, Egypt, paid",
        "uni_cohorts_ksa": "University cohorts, Saudi Arabia, paid",
        "uni_pilots": "University pilot cohorts, unpaid",
        "uni_licences": "Faculty licences",
        "school_cohorts": "School cohorts",
        "school_licences": "School licences",
        "tvet_cohorts": "Technical and vocational cohorts",
        "initiative_seats": "National initiative seats",
        "tot_cohorts": "Train the trainer cohorts",
        "ngo_programmes": "NGO and donor programmes",
        "competitions_self": "Competitions, self funded",
        "competitions_sponsored": "Competitions, sponsored",
        "corporates": "Corporate and CSR programmes",
        "podcast_seasons": "Mission Zero seasons carrying a sponsor",
    }
    for key, label in plan_labels.items():
        for y in range(1, 6):
            expected = sum(P.Q[key][(y - 1) * 4:y * 4])
            check(f"volume {key} year {y} ties to the plan",
                  expected, REV.y(label, y), tol=0.001)

    # --------------------------------------------------- revenue by line
    price = {
        "University cohort programmes": "University cohort, all in programme fee",
        "Faculty licences": "Faculty licence, annual",
        "School cohorts": "School cohort",
        "School licences": "School licence, annual",
        "Technical and vocational cohorts": "Technical and vocational cohort",
        "National initiative seats": "National initiative seat",
        "Train the trainer cohorts": "Train the trainer cohort",
        "NGO and donor programmes": "NGO or donor programme",
        "Sponsored competitions": "Sponsored competition season",
        "Corporate and CSR programmes": "Corporate or CSR programme",
        "Mission Zero sponsorship": "Mission Zero season sponsorship",
    }
    vol_for = {
        "University cohort programmes": None,
        "Faculty licences": "Faculty licences",
        "School cohorts": "School cohorts",
        "School licences": "School licences",
        "Technical and vocational cohorts": "Technical and vocational cohorts",
        "National initiative seats": "National initiative seats",
        "Train the trainer cohorts": "Train the trainer cohorts",
        "NGO and donor programmes": "NGO and donor programmes",
        "Sponsored competitions": "Competitions, sponsored",
        "Corporate and CSR programmes": "Corporate and CSR programmes",
        "Mission Zero sponsorship": "Mission Zero seasons carrying a sponsor",
    }
    for y in range(1, 6):
        p_col = 2 + y
        for stream, pl in price.items():
            pr = ass(pl, p_col)
            if stream == "University cohort programmes":
                vol = (sum(P.Q["uni_cohorts_eg"][(y - 1) * 4:y * 4])
                       + sum(P.Q["uni_cohorts_ksa"][(y - 1) * 4:y * 4]))
            else:
                vol = REV.y(vol_for[stream], y)
            check(f"{stream} year {y} equals volume times price",
                  vol * pr, REV.y(stream, y, nth=1 if stream in
                                  ("Faculty licences", "School cohorts",
                                   "School licences",
                                   "Technical and vocational cohorts",
                                   "National initiative seats",
                                   "Train the trainer cohorts",
                                   "NGO and donor programmes",
                                   "Corporate and CSR programmes") else 0))

    # ------------------------------------------------- revenue cross foot
    for y in range(1, 6):
        inst = sum(REV.y(s, y, nth=1 if s in
                         ("Faculty licences", "School cohorts",
                          "School licences",
                          "Technical and vocational cohorts",
                          "National initiative seats",
                          "Train the trainer cohorts",
                          "NGO and donor programmes",
                          "Corporate and CSR programmes") else 0)
                   for s in price) + REV.y("Programme and incubator seats", y,
                                           nth=1)
        check(f"institutional revenue year {y} is the sum of its lines",
              inst, REV.y("Institutional and sponsor revenue", y))
        check(f"total revenue year {y} is institutional plus subscriptions",
              inst + REV.y("Subscription revenue recognised", y),
              TOT.y("Total revenue", y))
        check(f"income statement revenue year {y} matches Total Revenue",
              TOT.y("Total revenue", y), IS.c("Total revenue", 2 + y))

    # ------------------------------------------------------ the funnel
    for y in (1, 3, 5):
        for m in (F.months_of(y)[0], F.months_of(y)[-1]):
            spend = ACQ.m("Short video, TikTok, Instagram, Facebook", m)
            cpc = ass("Short video, TikTok, Instagram, Facebook", 8)
            org = ass("Short video, TikTok, Instagram, Facebook", 9)
            ref = ass("Short video, TikTok, Instagram, Facebook", 10)
            s1 = ass("Short video, TikTok, Instagram, Facebook", 11)
            paid = spend / cpc if cpc else 0
            check(f"short video paid clicks, month {m}", paid,
                  ACQ.m("Paid clicks", m), tol=0.01)
            check(f"short video total clicks, month {m}",
                  paid * (1 + org + ref),
                  ACQ.m("Total clicks after organic and referral", m),
                  tol=0.01)
            check(f"short video leads, month {m}", paid * (1 + org + ref) * s1,
                  ACQ.m("Stage 1, leads", m), tol=0.01)
            s2 = ass("Stage 2, lead to registered account", 2 + y)
            check(f"registered from online leads, month {m}",
                  ACQ.m("Total online leads", m) * s2,
                  ACQ.m("Registered accounts from online leads, stage 2", m),
                  tol=0.01)
            s3i = ass("Stage 3, registered to paid, reached through an "
                      "institution", 2 + y)
            s3o = ass("Stage 3, registered to paid, reached online", 2 + y)
            check(f"new paid accounts, month {m}",
                  ACQ.m("Registered accounts from institutions", m) * s3i
                  + (ACQ.m("Registered accounts from online leads, stage 2", m)
                     + ACQ.m("Registered accounts from offline sign ups", m))
                  * s3o,
                  ACQ.m("Stage 3, new paid accounts", m), tol=0.01)

    # ------------------------------------------------------ direct cost
    for y in range(1, 6):
        rate = ass("Facilitator day rate", 2 + y)
        per = ass("Facilitators per cohort")
        days = ass("Delivery days, university cohort")
        uni = (REV.y("University cohorts, Egypt, paid", y)
               + REV.y("University cohorts, Saudi Arabia, paid", y)
               + REV.y("University pilot cohorts, unpaid", y))
        check(f"university facilitation year {y}", uni * rate * per * days,
              EXP.y("Facilitation, university cohorts", y))
        pd_share = ass("Party 1, share of cohorts delivered by certified "
                       "partners", 2 + y)
        pd_rate = ass("Party 1, the partner's share of that cohort fee")
        base = (REV.y("University cohort programmes", y)
                + REV.y("School cohorts", y, nth=1)
                + REV.y("Technical and vocational cohorts", y, nth=1))
        check(f"delivery partner revenue share year {y}",
              base * pd_share * pd_rate,
              EXP.y("Revenue share, certified delivery partners", y))

    # --------------------------------------------------------- payroll
    burden = PAY.v.cell(PAY.row("Employer burden on gross salary"), 3).value
    for y in range(1, 6):
        expected = sum(r[3][y - 1] * r[4][y - 1] for r in P.ROLES) \
            * 12 * (1 + burden)
        check(f"payroll year {y} equals establishment times salary",
              expected, PAY.y("Total payroll", y))
        check(f"headcount year {y} equals the plan",
              sum(r[4][y - 1] for r in P.ROLES),
              PAY.y("Total headcount", y), tol=0.001)
        check(f"headcount year {y} is within the cap",
              min(sum(r[4][y - 1] for r in P.ROLES), P.HEADCOUNT_CAP[y - 1]),
              PAY.y("Total headcount", y), tol=0.001)

    # ------------------------------------------------- profit cross foot
    for y in range(1, 6):
        rev = TOT.y("Total revenue", y)
        cogs = EXP.y("Total direct cost", y)
        gp = EXP.y("Gross profit", y)
        opex = EXP.y("Operating cost before depreciation", y)
        dna = EXP.y("Depreciation and amortisation", y)
        check(f"gross profit year {y} is revenue less direct cost",
              rev - cogs, gp)
        check(f"EBITDA year {y} is gross profit less operating cost",
              gp - opex, EXP.y("EBITDA", y))
        check(f"EBIT year {y} is EBITDA less depreciation",
              EXP.y("EBITDA", y) - dna, EXP.y("EBIT, operating profit", y))
        check(f"total expenses year {y} is direct plus operating",
              cogs + opex + dna, EXP.y("Total expenses, direct and operating",
                                       y))
        tax = IS.c("Corporate income tax", 2 + y)
        check(f"net profit year {y} is profit before tax plus tax",
              IS.c("Profit before tax", 2 + y) + tax,
              IS.c("Net profit for the year", 2 + y))
        check(f"tax year {y} is the rate on taxable profit",
              -IS.c("Taxable profit", 2 + y) * ass("Corporate income tax rate"),
              tax)

    # ---------------------------------------------------- balance sheet
    for y in range(1, 6):
        col = 3 + y
        check(f"balance sheet balances, year {y}",
              0, BS.c("CHECK, total assets less liabilities and equity", col),
              tol=0.01)
        check(f"cash on the balance sheet equals the cash flow, year {y}",
              CF.m("Cash at the end of the month", F.months_of(y)[-1]),
              BS.c("Cash and cash equivalents", col))
        check(f"retained earnings year {y} is cumulative net profit",
              sum(IS.c("Net profit for the year", 2 + k)
                  for k in range(1, y + 1)),
              BS.c("Retained earnings", col))
        seed = ass("Seed round raised")
        seed_m = ass("Month the seed round lands")
        check(f"share capital year {y} is opening cash plus the seed",
              ass("Opening cash") + (seed if F.months_of(y)[-1] >= seed_m
                                     else 0),
              BS.c("Share capital", col))
        check(f"total assets year {y} is the sum of its lines",
              BS.c("Cash and cash equivalents", col)
              + BS.c("Trade receivables", col)
              + BS.c("Property and equipment, net", col)
              + BS.c("Intangible assets, net", col),
              BS.c("Total assets", col))

    # -------------------------------------------------------- cash flow
    for y in range(1, 6):
        for m in (F.months_of(y)[0], F.months_of(y)[-1]):
            opening = (ass("Opening cash") if m == 1
                       else CF.m("Cash at the end of the month", m - 1))
            check(f"closing cash month {m} is opening plus the movement",
                  opening + CF.m("Net movement in cash", m),
                  CF.m("Cash at the end of the month", m))
            check(f"net movement month {m} is the three activities",
                  CF.m("Net cash from operating activities", m)
                  + CF.m("Net cash used in investing activities", m)
                  + CF.m("Net cash from financing activities", m),
                  CF.m("Net movement in cash", m))

    # ------------------------------------------- go to market capacity
    # The volume plan is only credible if the activity behind it can be run
    # by the people on the establishment. These recompute that claim rather
    # than reading the sheet's own answer back.
    GTM = Sheet(wv, wf, "GTM Capacity and Timing")
    sample = (1, 6, 13, 24, 30, 42, 55, 60)

    for m in sample:
        total_fn = sum(GTM.m(fn, m) for fn in FUNCTIONS)
        check(f"headcount by function month {m} is the establishment",
              PAY.m("Total headcount", m), total_fn)

    for i, (key, label, function, own, fnd, amb, prt) in enumerate(
            P.ACTIVITY_CAPACITY):
        for m in sample:
            expected = (GTM.m(function, m) * A.c(label, 3, nth=1)
                        + GTM.m("Leadership", m) * A.c(label, 4, nth=1)
                        + GTM.m("Campus ambassadors on a stipend", m)
                        * A.c(label, 5, nth=1)
                        + GTM.m("Certified delivery partner organisations", m)
                        * A.c(label, 6, nth=1))
            check(f"capacity for {key} month {m} is the four sources of "
                  f"people times their quotas",
                  expected,
                  GTM.m("Capacity, the people who can run it", m, nth=i))
            check(f"activity planned for {key} month {m} is what the "
                  f"Acquisitions sheet runs",
                  ACQ.m(label, m), GTM.m(label + ", planned", m))

    for m in sample:
        expected = sum(ACQ.m("Meetings held", m, nth=i) for i in range(5))
        check(f"meetings the funnel generates month {m} is the five buyer "
              f"types added up",
              expected, GTM.m("Meetings the funnel generates", m))
        check(f"meeting capacity month {m} is heads times their quotas",
              PAY.m("Commercial heads", m)
              * ass("Meetings per commercial head per month")
              + GTM.m("Leadership", m)
              * ass("Meetings per founder per month")
              + GTM.m("Growth and Content", m)
              * ass("Meetings per growth and content head per month"),
              GTM.m("Meetings the team can hold", m))
        expected = ((REV.m("University cohorts, Egypt, paid", m)
                     + REV.m("University cohorts, Saudi Arabia, paid", m)
                     + REV.m("University pilot cohorts, unpaid", m))
                    * ass("Delivery days, university cohort")
                    + REV.m("School cohorts", m)
                    * ass("Delivery days, school cohort")
                    + REV.m("Technical and vocational cohorts", m)
                    * ass("Delivery days, technical cohort")
                    ) * ass("Facilitators per cohort")
        check(f"facilitator days needed month {m} is cohorts times days "
              f"times facilitators",
              expected, GTM.m("Facilitator days the cohort plan needs", m))
        check(f"facilitator days available month {m} is the pool times the "
              f"days one gives",
              GTM.m("Certified facilitators in the pool", m)
              * ass("Delivery days one certified facilitator gives per "
                    "month"),
              GTM.m("Facilitator days the certified pool gives", m))

    # The three capacity gates. Every one of them must read nil in every
    # month, or the plan is selling something nobody can do.
    for label in ("Activity the plan asks for beyond what the team can run",
                  "Meetings asked for beyond capacity",
                  "Delivery days needed beyond the pool"):
        worst = max(GTM.m(label, m) for m in range(1, 61))
        results.append((worst <= 0.001,
                        f"capacity gate, {label.lower()}, reads nil in every "
                        f"month", 0, worst))

    # Signature and first delivery are arithmetic on the cycle and the
    # mobilisation period, counted from a first day of activity in month one.
    for i, (bkey, mob, why) in enumerate(P.MOBILISATION):
        label = BUYER_LABELS[bkey]
        start = GTM.c(label, 2)
        cycle = GTM.c(label, 3)
        check(f"{bkey} signs in month one plus its sales cycle",
              start + cycle, GTM.c(label, 4))
        check(f"{bkey} mobilisation period is the assumption",
              mob, GTM.c(label, 5))
        check(f"{bkey} first delivery is signature plus mobilisation",
              GTM.c(label, 4) + mob, GTM.c(label, 6))
        for m in (20, 36, 52):
            check(f"{bkey} contracts ready in month {m} are the ones signed "
                  f"{mob} months earlier",
                  ACQ.m("Contracts signed, after the sales cycle", m - mob,
                        nth=i),
                  GTM.m(label + ", mobilised and ready to deliver", m))

    # ----------------------------------------------- assets never negative
    for y in range(1, 6):
        m = F.months_of(y)[-1]
        gross = CAP.m("Total gross assets", m)
        accum = DEP.m("Total accumulated depreciation", m)
        nbv = DEP.m("Total net book value", m)
        check(f"net book value year {y} is cost less accumulated",
              gross - accum, nbv)
        results.append((nbv >= -0.01,
                        f"net book value year {y} is not negative", 0, nbv))
        results.append((accum <= gross + 0.01,
                        f"accumulated depreciation year {y} does not exceed "
                        f"cost", gross, accum))

    failed = [r for r in results if not r[0]]
    print(f"checks run:    {len(results)}")
    print(f"checks passed: {len(results) - len(failed)}")
    if failed:
        print(f"\nFAILED {len(failed)}:")
        for ok, name, exp, act in failed[:40]:
            print(f"  {name}\n      expected {exp:,.4f}   workbook "
                  f"{act:,.4f}   difference {(act or 0) - (exp or 0):,.4f}")
        return 1
    print("\nAll checks passed: every figure recomputes from the cells it "
          "claims to come from.")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1
                  else "output/StartPad_Financial_Model.xlsx"))
