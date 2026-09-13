# -*- coding: utf-8 -*-
"""Build the two-page Techne partnership proposal .docx from competition/proposal.py."""
import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from competition.proposal import (BLOCKS, CONTACT, EMAIL_BODY, EMAIL_SUBJECT,  # noqa: E402
                                  META, SUBTITLE, TITLE)
from scripts._determinism import FIXED_TIMESTAMP, normalize_zip  # noqa: E402

INK = RGBColor(0x10, 0x1E, 0x2E)
SEA = RGBColor(0x0E, 0x7C, 0x86)
GREY = RGBColor(0x63, 0x73, 0x7F)


def para(doc, text="", size=10.5, bold=False, italic=False, color=INK,
         space_before=0, space_after=6, align=None, indent=0, line=1.2):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    if indent:
        pf.left_indent = Cm(indent)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return p


def rule(doc, color="101E2E", size="8", after=10):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(after)
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    p._p.get_or_add_pPr().append(pbdr)


def bullet(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("—  ")
    r.font.size = Pt(10.5)
    r.font.color.rgb = SEA
    r.font.bold = True
    r2 = p.add_run(text)
    r2.font.size = Pt(10.5)
    r2.font.color.rgb = INK


def build(out_path):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.font.color.rgb = INK
    for s in doc.sections:
        s.top_margin = Cm(1.9)
        s.bottom_margin = Cm(1.9)
        s.left_margin = Cm(2.1)
        s.right_margin = Cm(2.1)

    para(doc, "PROPOSAL  ·  PARTNERSHIP", size=8.5, bold=True, color=SEA, space_after=4)
    para(doc, TITLE, size=19, bold=True, color=INK, space_after=3, line=1.1)
    para(doc, SUBTITLE, size=11, italic=True, color=GREY, space_after=10, line=1.25)
    rule(doc)
    for label, value in META:
        p = para(doc, "", size=9.5, space_after=2)
        r = p.add_run(label + ": ")
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = GREY
        r2 = p.add_run(value)
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = INK
    rule(doc, color="D9DBD3", size="6", after=12)

    for block in BLOCKS:
        if block[0] == "h":
            para(doc, block[1], size=12.5, bold=True, color=SEA,
                 space_before=12, space_after=4)
        elif block[0] == "p":
            para(doc, block[1], size=10.5, space_after=7, line=1.25)
        elif block[0] == "b":
            for item in block[1]:
                bullet(doc, item)

    rule(doc, color="D9DBD3", size="6", after=8)
    for label, value in CONTACT:
        p = para(doc, "", size=9.5, space_after=2)
        r = p.add_run(label + "   ")
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = INK
        r2 = p.add_run(value)
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = GREY

    doc.add_page_break()
    para(doc, "COVERING EMAIL", size=8.5, bold=True, color=SEA, space_after=4)
    para(doc, "Send this, attach the two pages above", size=13, bold=True, space_after=8)
    rule(doc, color="D9DBD3", size="6")
    p = para(doc, "", size=10.5, space_after=10)
    r = p.add_run("Subject:  ")
    r.font.bold = True
    r.font.size = Pt(10.5)
    r2 = p.add_run(EMAIL_SUBJECT)
    r2.font.size = Pt(10.5)
    for line in EMAIL_BODY.strip().split("\n"):
        para(doc, line, size=10.5, space_after=0 if line.strip() else 6, line=1.25)

    core = doc.core_properties
    core.title = TITLE
    core.author = "Abdelrhman Soliman"
    core.created = FIXED_TIMESTAMP
    core.modified = FIXED_TIMESTAMP

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    doc.save(out_path)
    normalize_zip(out_path)
    return out_path


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else \
        "output/StartPad_Techne_Partnership_Proposal.docx"
    p = build(out)
    print("wrote", p, os.path.getsize(p), "bytes")
